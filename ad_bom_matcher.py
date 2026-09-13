"""BOM 物料库匹配工具：界面、规格识别与匹配。

作者：桀桀
联系作者：微信号 JJ-Linnnnn
"""

import csv
import hashlib
import json
import os
import re
import sys
import traceback
import tempfile
import uuid
import zipfile
import xml.etree.ElementTree as ET
from concurrent.futures import Future, ThreadPoolExecutor
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Any, Callable, Dict, List, Optional, Tuple

import tkinter as tk
import tkinter.font as tkfont
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageTk
import xlrd



APP_TITLE = "BOM 物料库匹配工具 · 开源版"
APP_AUTHOR = "桀桀"
AUTHOR_WECHAT = "JJ-Linnnnn"
WORKER_COUNT = max(2, min(8, os.cpu_count() or 2))
AUTO_MATCH_THRESHOLD = 72
AMBIGUITY_MARGIN = 6
SPEC_PARSER_VERSION = 4
APP_SETTINGS_DIR = "BOMMaterialMatcherOpenSource"
APP_SETTINGS_FILE = "settings.json"
DEFAULT_APP_SETTINGS = {
    "auto_load_last_library": False,
    "last_library_path": "",
    "custom_result_columns_enabled": False,
    "result_column_order": [],
    "result_column_titles": {},
}


def resource_path(filename: str) -> str:
    """Resolve bundled assets in both source and PyInstaller one-file runs."""
    bundle_root = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(bundle_root, filename)


def app_settings_path() -> str:
    root = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA") or os.path.expanduser("~")
    return os.path.join(root, APP_SETTINGS_DIR, APP_SETTINGS_FILE)


def load_app_settings() -> Dict[str, Any]:
    settings = dict(DEFAULT_APP_SETTINGS)
    try:
        with open(app_settings_path(), "r", encoding="utf-8") as settings_file:
            saved = json.load(settings_file)
    except (OSError, ValueError, json.JSONDecodeError):
        return settings
    if not isinstance(saved, dict):
        return settings
    settings["auto_load_last_library"] = bool(saved.get("auto_load_last_library", False))
    last_path = saved.get("last_library_path", "")
    settings["last_library_path"] = last_path if isinstance(last_path, str) else ""
    settings["custom_result_columns_enabled"] = bool(saved.get("custom_result_columns_enabled", False))
    saved_order = saved.get("result_column_order", [])
    settings["result_column_order"] = [item for item in saved_order if isinstance(item, str)] if isinstance(saved_order, list) else []
    saved_titles = saved.get("result_column_titles", {})
    settings["result_column_titles"] = (
        {str(key): value for key, value in saved_titles.items() if isinstance(key, str) and isinstance(value, str)}
        if isinstance(saved_titles, dict)
        else {}
    )
    return settings


def save_app_settings(settings: Dict[str, Any]) -> None:
    path = app_settings_path()
    directory = os.path.dirname(path)
    os.makedirs(directory, exist_ok=True)
    temp_path = os.path.join(directory, f".{APP_SETTINGS_FILE}.{uuid.uuid4().hex}.tmp")
    payload = {
        "auto_load_last_library": bool(settings.get("auto_load_last_library", False)),
        "last_library_path": str(settings.get("last_library_path", "")),
        "custom_result_columns_enabled": bool(settings.get("custom_result_columns_enabled", False)),
        "result_column_order": list(settings.get("result_column_order", [])),
        "result_column_titles": dict(settings.get("result_column_titles", {})),
    }
    try:
        with open(temp_path, "w", encoding="utf-8") as settings_file:
            json.dump(payload, settings_file, ensure_ascii=False, indent=2)
        os.replace(temp_path, path)
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass

OUTPUT_COLUMNS = [
    "层级",
    "工序",
    "物料编码",
    "物料名称",
    "规格",
    "用量",
    "单位",
    "位号",
    "点数",
]

RESULT_COLUMN_KEYS = ["确认"] + OUTPUT_COLUMNS + ["备注"]
CUSTOMIZABLE_RESULT_COLUMN_KEYS = OUTPUT_COLUMNS + ["备注"]
LEGACY_RESULT_COLUMN_KEY_MAP = {
    "（贴片1 后焊2）工序": "工序",
    "总用量": "用量",
    "备注（位号）": "位号",
    "工序工艺": "点数",
    "匹配分数": "备注",
}


def default_result_column_title(key: str) -> str:
    return key


def canonical_result_column_key(key: str) -> str:
    return LEGACY_RESULT_COLUMN_KEY_MAP.get(key, key)


def normalize_result_column_order(value: Any) -> List[str]:
    requested = value if isinstance(value, list) else []
    normalized: List[str] = []
    for key in requested:
        canonical_key = canonical_result_column_key(key) if isinstance(key, str) else ""
        if canonical_key in CUSTOMIZABLE_RESULT_COLUMN_KEYS and canonical_key not in normalized:
            normalized.append(canonical_key)
    normalized.extend(key for key in CUSTOMIZABLE_RESULT_COLUMN_KEYS if key not in normalized)
    return normalized


def normalize_result_column_titles(value: Any) -> Dict[str, str]:
    if not isinstance(value, dict):
        return {}
    normalized: Dict[str, str] = {}
    for key, title in value.items():
        canonical_key = canonical_result_column_key(key) if isinstance(key, str) else ""
        cleaned_title = clean_cell(title).strip() if isinstance(title, str) else ""
        if canonical_key in CUSTOMIZABLE_RESULT_COLUMN_KEYS and cleaned_title:
            normalized[canonical_key] = cleaned_title
    return normalized

LIB_ALIASES = {
    "code": ["物料编码", "产品编码", "产品编号", "物料代码", "物料号", "料号", "品号", "编码", "编号", "存货编码", "零件号", "part number", "part no", "material code", "item code", "item no", "p/n", "pn"],
    "name": ["物料名称", "名称", "品名", "材料名称", "description", "part name", "name"],
    "spec": ["规格", "规格型号", "型号规格", "型号", "规格描述", "参数", "封装规格", "spec", "specification", "model", "part spec", "comment", "value"],
    "unit": ["单位", "计量单位", "unit", "uom"],
}

CODE_HEADER_EXCLUDES = ["序号", "行号", "序列号", "流水号", "no", "num", "index", "id"]

BOM_ALIASES = {
    "spec": [
        "comment",
        "value",
        "规格",
        "规格型号",
        "型号",
        "描述",
        "description",
        "libref",
        "component",
        "part type",
        "参数",
    ],
    "qty": ["quantity", "qty", "数量", "用量", "总用量", "count"],
    "designator": ["designator", "designators", "reference", "references", "位号", "元件位号", "refdes", "ref"],
    "footprint": ["footprint", "pcb footprint", "封装", "package", "封装名称"],
}

POST_WELD_KEYWORDS = [
    "dip",
    "th",
    "through",
    "插件",
    "后焊",
    "直插",
    "连接器",
    "端子",
    "排针",
    "排母",
    "座",
    "开关",
    "button",
    "switch",
    "connector",
    "header",
    "terminal",
    "usb",
    "dc",
    "jack",
    "con",
    "jst",
    "screw",
]

CHIP_PACKAGES = ["0201", "0402", "0603", "0805", "1206", "1210", "1812", "2010", "2512"]
METRIC_PACKAGE_ALIASES = {
    "0603m": "0201",
    "1005": "0402",
    "1608": "0603",
    "2012": "0805",
    "3216": "1206",
    "3225": "1210",
    "4532": "1812",
    "5025": "2010",
    "6332": "2512",
}
CAPACITOR_VOLTAGE_CODES = {
    "0j": 6.3,
    "1a": 10.0,
    "1c": 16.0,
    "1e": 25.0,
    "1h": 50.0,
    "2a": 100.0,
    "2d": 200.0,
    "2e": 250.0,
    "2w": 450.0,
}

# Canonical package families. Manufacturer and EDA-library spellings vary,
# but the physically relevant identity is still family + pins + body size.
PACKAGE_FAMILY_ALIASES = {
    "qfn": "qfn", "vqfn": "qfn", "wqfn": "qfn", "hvqfn": "qfn",
    "pqfn": "qfn", "tqfn": "qfn", "uqfn": "qfn", "xqfn": "qfn",
    "mlf": "qfn", "mlpq": "qfn", "lfcsp": "qfn",
    "dfn": "dfn", "wdfn": "dfn", "tdfn": "dfn", "udfn": "dfn", "pdfn": "dfn",
    "x2dfn": "dfn", "son": "dfn", "vson": "dfn", "wson": "dfn",
    "lga": "lga", "bga": "bga", "fbga": "bga", "tfbga": "bga",
    "ucbga": "bga", "csp": "csp", "wlcsp": "csp", "dsbga": "csp",
    "qfp": "qfp", "lqfp": "lqfp", "tqfp": "tqfp", "vqfp": "qfp",
    "pqfp": "qfp", "mqfp": "qfp",
    "sop": "sop", "so": "sop", "soic": "sop", "hsop": "hsop",
    "psop": "psop", "ssop": "ssop", "tssop": "tssop",
    "msop": "msop", "vsop": "vsop", "tsop": "tsop",
    "dip": "dip", "pdip": "dip", "sdip": "sdip", "sip": "sip",
    "plcc": "plcc", "clcc": "clcc", "lcc": "lcc",
}
PACKAGE_FAMILY_PATTERN = "|".join(
    sorted((re.escape(item) for item in PACKAGE_FAMILY_ALIASES), key=len, reverse=True)
)

COMPONENT_KIND_KEYWORDS = (
    ("fuse", ("保险丝", "保险管", "熔断器", "resettable fuse", "polyfuse", "pptc", "fuse")),
    ("connector", ("连接器", "接插件", "排针", "排母", "端子", "插座", "母座", "公座", "插头", "接头", "接口", "type-c", "type c", "usb-c", "usb c", "micro usb", "usb", "connector", "header", "socket", "terminal", "jack", "hdr", "jst")),
    ("led", ("发光二极管", "led")),
    ("tvs", ("瞬态抑制", "浪涌保护", "esd protection", "tvs")),
    ("mosfet", ("场效应管", "mosfet", "mos-f", "mos-n", "mos-p", "nmos", "pmos")),
    ("transistor", ("三极管", "晶体管", "transistor", "bjt", "npn", "pnp")),
    ("diode", ("肖特基", "整流二极管", "稳压二极管", "二极管", "schottky", "rectifier", "zener", "diode")),
    ("crystal", ("晶振", "晶体", "oscillator", "crystal")),
    ("relay", ("继电器", "relay")),
    ("switch", ("轻触开关", "拨动开关", "按键开关", "按键", "switch", "button", "power key")),
    ("transformer", ("变压器", "transformer")),
    ("buzzer", ("蜂鸣器", "扬声器", "buzzer", "speaker")),
    ("optocoupler", ("光耦", "光电耦合", "optocoupler", "optoisolator")),
    ("thermistor", ("热敏电阻", "thermistor", "ntc", "ptc")),
    ("varistor", ("压敏电阻", "varistor", "mov")),
    ("potentiometer", ("电位器", "可调电阻", "potentiometer", "trimmer")),
    ("battery", ("电池", "battery")),
    ("antenna", ("天线", "antenna")),
    ("motor", ("电机", "马达", "motor")),
    ("sensor", ("传感器", "sensor")),
    ("module", ("模块", "module")),
    ("testpoint", ("测试点", "test point", "testpoint")),
    ("inductor", ("电感", "磁珠", "inductor", "choke", "ferrite bead")),
    ("capacitor", ("电容", "capacitor", "mlcc")),
    ("resistor", ("电阻", "resistor", "resistance")),
    ("ic", ("集成电路", "芯片", "integrated circuit")),
)


@dataclass
class MaterialItem:
    code: str = ""
    name: str = ""
    spec: str = ""
    unit: str = "pcs"
    source: Dict[str, str] = field(default_factory=dict)
    signature: Optional["ComponentSignature"] = None
    normalized_searchable: str = ""
    identifiers: frozenset[str] = field(default_factory=frozenset)
    analysis_cache_hit: bool = False

    @property
    def searchable(self) -> str:
        source_text = " ".join(clean_cell(value) for value in self.source.values())
        return " ".join([self.code, self.name, self.spec, source_text]).strip()


@dataclass
class BomItem:
    original: Dict[str, str]
    spec: str
    qty: str
    designator: str
    footprint: str


@dataclass(frozen=True)
class ComponentSignature:
    kind: Optional[str] = None
    value_key: Optional[str] = None
    package: Optional[str] = None
    package_family: Optional[str] = None
    pin_count: Optional[int] = None
    package_width_mm: Optional[float] = None
    package_height_mm: Optional[float] = None
    tolerance: Optional[str] = None
    tolerance_percent: Optional[float] = None
    voltage_v: Optional[float] = None
    current_a: Optional[float] = None
    dielectric: Optional[str] = None
    subtype: Optional[str] = None


def material_analysis_cache_path(source_path: str) -> str:
    normalized_path = os.path.normcase(os.path.abspath(source_path))
    digest = hashlib.sha256(normalized_path.encode("utf-8", errors="ignore")).hexdigest()
    return os.path.join(os.path.dirname(app_settings_path()), "spec_cache", f"{digest}.json")


def material_analysis_key(material: MaterialItem) -> str:
    payload = {
        "code": clean_cell(material.code),
        "name": clean_cell(material.name),
        "spec": clean_cell(material.spec),
        "unit": clean_cell(material.unit),
        "source": {str(key): clean_cell(value) for key, value in sorted(material.source.items())},
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8", errors="ignore")).hexdigest()


def signature_to_cache_dict(signature: ComponentSignature) -> Dict[str, Any]:
    return {
        "kind": signature.kind,
        "value_key": signature.value_key,
        "package": signature.package,
        "package_family": signature.package_family,
        "pin_count": signature.pin_count,
        "package_width_mm": signature.package_width_mm,
        "package_height_mm": signature.package_height_mm,
        "tolerance": signature.tolerance,
        "tolerance_percent": signature.tolerance_percent,
        "voltage_v": signature.voltage_v,
        "current_a": signature.current_a,
        "dielectric": signature.dielectric,
        "subtype": signature.subtype,
    }


def signature_from_cache_dict(value: Any) -> Optional[ComponentSignature]:
    if not isinstance(value, dict):
        return None
    return ComponentSignature(
        kind=value.get("kind") if isinstance(value.get("kind"), str) else None,
        value_key=value.get("value_key") if isinstance(value.get("value_key"), str) else None,
        package=value.get("package") if isinstance(value.get("package"), str) else None,
        package_family=value.get("package_family") if isinstance(value.get("package_family"), str) else None,
        pin_count=int(value["pin_count"]) if isinstance(value.get("pin_count"), (int, float)) else None,
        package_width_mm=float(value["package_width_mm"]) if isinstance(value.get("package_width_mm"), (int, float)) else None,
        package_height_mm=float(value["package_height_mm"]) if isinstance(value.get("package_height_mm"), (int, float)) else None,
        tolerance=value.get("tolerance") if isinstance(value.get("tolerance"), str) else None,
        tolerance_percent=float(value["tolerance_percent"]) if isinstance(value.get("tolerance_percent"), (int, float)) else None,
        voltage_v=float(value["voltage_v"]) if isinstance(value.get("voltage_v"), (int, float)) else None,
        current_a=float(value["current_a"]) if isinstance(value.get("current_a"), (int, float)) else None,
        dielectric=value.get("dielectric") if isinstance(value.get("dielectric"), str) else None,
        subtype=value.get("subtype") if isinstance(value.get("subtype"), str) else None,
    )


def load_material_analysis_cache(source_path: str) -> Dict[str, Tuple[ComponentSignature, frozenset[str]]]:
    try:
        source_stat = os.stat(source_path)
        with open(material_analysis_cache_path(source_path), "r", encoding="utf-8") as cache_file:
            payload = json.load(cache_file)
    except (OSError, ValueError, json.JSONDecodeError):
        return {}
    if not isinstance(payload, dict):
        return {}
    if payload.get("parser_version") != SPEC_PARSER_VERSION:
        return {}
    if payload.get("source_size") != source_stat.st_size or payload.get("source_mtime_ns") != source_stat.st_mtime_ns:
        return {}
    entries = payload.get("entries")
    if not isinstance(entries, dict):
        return {}
    result: Dict[str, Tuple[ComponentSignature, frozenset[str]]] = {}
    for key, cached in entries.items():
        if not isinstance(key, str) or not isinstance(cached, dict):
            continue
        signature = signature_from_cache_dict(cached.get("signature"))
        identifiers = cached.get("identifiers", [])
        if signature is None or not isinstance(identifiers, list):
            continue
        result[key] = (signature, frozenset(item for item in identifiers if isinstance(item, str)))
    return result


def save_material_analysis_cache(source_path: str, materials: List[MaterialItem]) -> None:
    source_stat = os.stat(source_path)
    cache_path = material_analysis_cache_path(source_path)
    cache_directory = os.path.dirname(cache_path)
    os.makedirs(cache_directory, exist_ok=True)
    entries = {
        material_analysis_key(material): {
            "signature": signature_to_cache_dict(material.signature),
            "identifiers": sorted(material.identifiers),
        }
        for material in materials
        if material.signature is not None
    }
    payload = {
        "parser_version": SPEC_PARSER_VERSION,
        "source_path": os.path.abspath(source_path),
        "source_size": source_stat.st_size,
        "source_mtime_ns": source_stat.st_mtime_ns,
        "entries": entries,
    }
    temp_path = f"{cache_path}.{uuid.uuid4().hex}.tmp"
    try:
        with open(temp_path, "w", encoding="utf-8") as cache_file:
            json.dump(payload, cache_file, ensure_ascii=False, separators=(",", ":"))
        os.replace(temp_path, cache_path)
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


def resource_path(relative: str) -> str:
    if hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, relative)
    return os.path.join(os.path.abspath("."), relative)


def clean_header(value) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\n", " ").replace("\r", " ").strip()
    return re.sub(r"\s+", " ", text)


def clean_cell(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and re.fullmatch(r"-?\d+\.0", text):
        text = text[:-2]
    return text


def normalize_key(value: str) -> str:
    return re.sub(r"[\s_()（）【】\[\]{}:：;；/\\\-]+", "", str(value).strip().lower())


def normalize_text(value: str) -> str:
    text = str(value or "").lower()
    text = re.sub(r"(?<=\d)\s*([km])\s*(Ω|ω|ohm)\b", r"\1", text, flags=re.I)
    text = re.sub(r"(?<=\d)\s*(Ω|ω|ohm)\b", "r", text, flags=re.I)
    replacements = {
        "μ": "u",
        "µ": "u",
        "Ω": "r",
        "ω": "r",
        "ohm": "r",
        "％": "%",
        "，": ",",
        "。": ".",
        "（": "(",
        "）": ")",
        "　": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"(?<=\d)\s+(?=[a-z%])", "", text)
    text = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff.%+/\-]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def alias_match(headers: List[str], aliases: List[str]) -> Optional[str]:
    normalized = {normalize_key(header): header for header in headers}
    for alias in aliases:
        key = normalize_key(alias)
        if key in normalized:
            return normalized[key]
    for header in headers:
        header_key = normalize_key(header)
        for alias in aliases:
            alias_key = normalize_key(alias)
            if alias_key and alias_key in header_key:
                return header
    return None


def score_header(header: str, aliases: List[str], excludes: Optional[List[str]] = None) -> int:
    key = normalize_key(header)
    if not key:
        return -100
    for exclude in excludes or []:
        ex_key = normalize_key(exclude)
        if key == ex_key or key.endswith(ex_key):
            return -50
    best = 0
    for alias in aliases:
        alias_key = normalize_key(alias)
        if not alias_key:
            continue
        if key == alias_key:
            best = max(best, 100)
        elif alias_key in key:
            best = max(best, 80)
        elif key in alias_key:
            best = max(best, 55)
    return best


def best_header_match(headers: List[str], aliases: List[str], excludes: Optional[List[str]] = None) -> Optional[str]:
    ranked = sorted(((score_header(header, aliases, excludes), -index, header) for index, header in enumerate(headers)), reverse=True)
    if ranked and ranked[0][0] >= 55:
        return ranked[0][2]
    return None


def looks_like_material_code(value: str) -> bool:
    text = clean_cell(value)
    if not text or re.fullmatch(r"\d{1,4}", text):
        return False
    return bool(re.search(r"[A-Za-z]", text) and re.search(r"\d", text))


def infer_material_code_header(headers: List[str], rows: List[Dict[str, str]]) -> Optional[str]:
    best_header = None
    best_score = 0
    for header in headers:
        if score_header(header, [], CODE_HEADER_EXCLUDES) < 0:
            continue
        values = [row.get(header, "") for row in rows[:80]]
        non_empty = [value for value in values if clean_cell(value)]
        if not non_empty:
            continue
        ratio = sum(1 for value in non_empty if looks_like_material_code(value)) / len(non_empty)
        avg_len = sum(len(clean_cell(value)) for value in non_empty) / len(non_empty)
        score = ratio * 100 + min(avg_len, 20)
        if score > best_score:
            best_score = score
            best_header = header
    return best_header if best_score >= 45 else None


def first_existing(row: Dict[str, str], headers: List[str], aliases: List[str]) -> str:
    matched = alias_match(headers, aliases)
    if matched:
        return row.get(matched, "")
    for header in headers:
        value = row.get(header, "")
        if value:
            return value
    return ""


def first_existing_non_index(row: Dict[str, str], headers: List[str], aliases: List[str]) -> str:
    matched = alias_match(headers, aliases)
    if matched:
        return row.get(matched, "")
    for header in headers:
        if score_header(header, [], CODE_HEADER_EXCLUDES) < 0:
            continue
        value = row.get(header, "")
        if value:
            return value
    return ""


def detect_header_index(raw_rows: List[List[str]]) -> int:
    best_index = 0
    best_score = -1
    for index, row in enumerate(raw_rows[:20]):
        non_empty = [clean_cell(value) for value in row if clean_cell(value)]
        if not non_empty:
            continue
        joined = normalize_key(" ".join(non_empty))
        keyword_score = 0
        for keyword in ["designator", "quantity", "comment", "物料", "规格", "数量", "位号", "编码"]:
            if normalize_key(keyword) in joined:
                keyword_score += 5
        score = len(non_empty) + keyword_score
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def make_unique_headers(headers: List[str]) -> List[str]:
    used: Dict[str, int] = {}
    result = []
    for index, header in enumerate(headers, start=1):
        base = clean_header(header) or f"列{index}"
        count = used.get(base, 0)
        used[base] = count + 1
        result.append(base if count == 0 else f"{base}_{count + 1}")
    return result


def read_xlsx_rows(path: str) -> List[List[str]]:
    # Read cell values directly from the XLSX package so unsupported style fields
    # from some exported spreadsheets, such as xcid, cannot block import.
    with zipfile.ZipFile(path, "r") as workbook_zip:
        shared_strings = read_shared_strings(workbook_zip)
        sheet_path = first_sheet_path(workbook_zip)
        sheet_xml = workbook_zip.read(sheet_path)
    return parse_sheet_xml(sheet_xml, shared_strings)


def first_sheet_path(workbook_zip: zipfile.ZipFile) -> str:
    default_path = "xl/worksheets/sheet1.xml"
    names = set(workbook_zip.namelist())
    if "xl/workbook.xml" not in names or "xl/_rels/workbook.xml.rels" not in names:
        return default_path

    workbook_root = ET.fromstring(workbook_zip.read("xl/workbook.xml"))
    rels_root = ET.fromstring(workbook_zip.read("xl/_rels/workbook.xml.rels"))
    rel_map = {
        rel.attrib.get("Id"): rel.attrib.get("Target", "")
        for rel in rels_root
        if rel.attrib.get("Id")
    }
    for sheet in workbook_root.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
        rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = rel_map.get(rel_id, "")
        if target:
            target = target.replace("\\", "/")
            if target.startswith("/"):
                return target.lstrip("/")
            if target.startswith("xl/"):
                return target
            return "xl/" + target.lstrip("/")
    return default_path


def read_shared_strings(workbook_zip: zipfile.ZipFile) -> List[str]:
    if "xl/sharedStrings.xml" not in workbook_zip.namelist():
        return []
    root = ET.fromstring(workbook_zip.read("xl/sharedStrings.xml"))
    values: List[str] = []
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    for item in root.findall(f"{namespace}si"):
        texts = [node.text or "" for node in item.findall(f".//{namespace}t")]
        values.append("".join(texts))
    return values


def parse_sheet_xml(sheet_xml: bytes, shared_strings: List[str]) -> List[List[str]]:
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    root = ET.fromstring(sheet_xml)
    rows: List[List[str]] = []
    for row in root.findall(f".//{namespace}sheetData/{namespace}row"):
        values: Dict[int, str] = {}
        max_col = 0
        for cell in row.findall(f"{namespace}c"):
            ref = cell.attrib.get("r", "")
            col_index = column_index_from_ref(ref) if ref else max_col + 1
            max_col = max(max_col, col_index)
            values[col_index] = cell_value(cell, shared_strings)
        rows.append([values.get(col, "") for col in range(1, max_col + 1)])
    return rows


def column_index_from_ref(ref: str) -> int:
    letters = re.match(r"[A-Z]+", ref.upper())
    if not letters:
        return 1
    index = 0
    for char in letters.group(0):
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


def cell_value(cell: ET.Element, shared_strings: List[str]) -> str:
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    cell_type = cell.attrib.get("t", "")
    if cell_type == "inlineStr":
        texts = [node.text or "" for node in cell.findall(f".//{namespace}t")]
        return clean_cell("".join(texts))

    value_node = cell.find(f"{namespace}v")
    if value_node is None:
        value_node = cell.find(f"{namespace}is/{namespace}t")
    value = value_node.text if value_node is not None else ""
    if cell_type == "s":
        try:
            return clean_cell(shared_strings[int(value)])
        except (ValueError, IndexError):
            return ""
    if cell_type == "b":
        return "TRUE" if value == "1" else "FALSE"
    return clean_cell(value)


def read_xlsx_rows_with_openpyxl(path: str) -> List[List[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    raw_rows = [[clean_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return raw_rows


def remove_xcid_from_xlsx(source_path: str, target_path: str) -> None:
    xcid_pattern = re.compile(rb'\s+xcid="[^"]*"')
    with zipfile.ZipFile(source_path, "r") as source_zip:
        with zipfile.ZipFile(target_path, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
            for item in source_zip.infolist():
                content = source_zip.read(item.filename)
                if item.filename == "xl/styles.xml":
                    content = xcid_pattern.sub(b"", content)
                target_zip.writestr(item, content)


def read_table(path: str) -> Tuple[List[str], List[Dict[str, str]]]:
    ext = os.path.splitext(path)[1].lower()
    if ext in [".xlsx", ".xlsm", ".xltx", ".xltm"]:
        raw_rows = read_xlsx_rows(path)
    elif ext == ".xls":
        workbook = xlrd.open_workbook(path)
        sheet = workbook.sheet_by_index(0)
        raw_rows = [[clean_cell(sheet.cell_value(row_index, col_index)) for col_index in range(sheet.ncols)] for row_index in range(sheet.nrows)]
    elif ext == ".csv":
        raw_rows = read_csv_rows(path)
    else:
        raise ValueError("仅支持 .xlsx、.xlsm、.xls 和 .csv 文件")

    raw_rows = [trim_trailing_empty(row) for row in raw_rows]
    raw_rows = [row for row in raw_rows if any(clean_cell(value) for value in row)]
    if not raw_rows:
        return [], []

    header_index = detect_header_index(raw_rows)
    headers = make_unique_headers(raw_rows[header_index])
    width = len(headers)
    rows: List[Dict[str, str]] = []
    for raw in raw_rows[header_index + 1 :]:
        padded = raw + [""] * max(0, width - len(raw))
        record = {headers[i]: clean_cell(padded[i]) if i < len(padded) else "" for i in range(width)}
        if any(record.values()):
            rows.append(record)
    return headers, rows


def trim_trailing_empty(row: List[str]) -> List[str]:
    result = list(row)
    while result and not clean_cell(result[-1]):
        result.pop()
    return result


def read_csv_rows(path: str) -> List[List[str]]:
    encodings = ["utf-8-sig", "gbk", "gb2312", "utf-16"]
    last_error = None
    for encoding in encodings:
        try:
            with open(path, "r", newline="", encoding=encoding) as handle:
                sample = handle.read(4096)
                handle.seek(0)
                dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
                return [[clean_cell(value) for value in row] for row in csv.reader(handle, dialect)]
        except Exception as exc:
            last_error = exc
    raise ValueError(f"CSV 文件读取失败：{last_error}")


def parse_materials(
    headers: List[str],
    rows: List[Dict[str, str]],
    analysis_cache: Optional[Dict[str, Tuple[ComponentSignature, frozenset[str]]]] = None,
) -> List[MaterialItem]:
    code_header = best_header_match(headers, LIB_ALIASES["code"], CODE_HEADER_EXCLUDES) or infer_material_code_header(headers, rows)
    name_header = best_header_match(headers, LIB_ALIASES["name"])
    spec_header = best_header_match(headers, LIB_ALIASES["spec"])
    unit_header = best_header_match(headers, LIB_ALIASES["unit"])

    materials: List[MaterialItem] = []
    for row in rows:
        code = row.get(code_header, "") if code_header else ""
        name = row.get(name_header, "") if name_header else ""
        spec = row.get(spec_header, "") if spec_header else ""
        unit = row.get(unit_header, "") if unit_header else ""

        if not spec:
            spec = first_existing_non_index(row, headers, LIB_ALIASES["spec"])
        spec = choose_best_material_spec(row, headers, spec)
        if not name and len(headers) >= 2:
            name = row.get(headers[1], "")
        if code or name or spec:
            material = MaterialItem(code=code, name=name, spec=spec, unit=unit or "pcs", source=row)
            apply_material_analysis(material, analysis_cache)
            materials.append(material)
    return materials


def material_from_row(
    row: Dict[str, str],
    headers: List[str],
    code_header: Optional[str],
    name_header: Optional[str],
    spec_header: Optional[str],
    unit_header: Optional[str],
    analysis_cache: Optional[Dict[str, Tuple[ComponentSignature, frozenset[str]]]] = None,
) -> Optional[MaterialItem]:
    code = row.get(code_header, "") if code_header else ""
    name = row.get(name_header, "") if name_header else ""
    spec = row.get(spec_header, "") if spec_header else ""
    unit = row.get(unit_header, "") if unit_header else ""

    if not spec:
        spec = first_existing_non_index(row, headers, LIB_ALIASES["spec"])
    spec = choose_best_material_spec(row, headers, spec)
    if not name and len(headers) >= 2:
        name = row.get(headers[1], "")
    if not (code or name or spec):
        return None
    material = MaterialItem(code=code, name=name, spec=spec, unit=unit or "pcs", source=row)
    apply_material_analysis(material, analysis_cache)
    return material


def material_chunk_from_rows(
    chunk: List[Dict[str, str]],
    headers: List[str],
    code_header: Optional[str],
    name_header: Optional[str],
    spec_header: Optional[str],
    unit_header: Optional[str],
    analysis_cache: Optional[Dict[str, Tuple[ComponentSignature, frozenset[str]]]] = None,
) -> List[MaterialItem]:
    materials: List[MaterialItem] = []
    for row in chunk:
        material = material_from_row(row, headers, code_header, name_header, spec_header, unit_header, analysis_cache)
        if material is not None:
            materials.append(material)
    return materials


def parse_materials_fast(
    headers: List[str],
    rows: List[Dict[str, str]],
    analysis_cache: Optional[Dict[str, Tuple[ComponentSignature, frozenset[str]]]] = None,
) -> List[MaterialItem]:
    code_header = best_header_match(headers, LIB_ALIASES["code"], CODE_HEADER_EXCLUDES) or infer_material_code_header(headers, rows)
    name_header = best_header_match(headers, LIB_ALIASES["name"])
    spec_header = best_header_match(headers, LIB_ALIASES["spec"])
    unit_header = best_header_match(headers, LIB_ALIASES["unit"])

    if len(rows) < 800:
        return parse_materials(headers, rows, analysis_cache)

    chunk_size = max(300, min(1200, len(rows) // max(1, WORKER_COUNT * 3)))
    chunks = [rows[index : index + chunk_size] for index in range(0, len(rows), chunk_size)]
    materials: List[MaterialItem] = []
    with ThreadPoolExecutor(max_workers=WORKER_COUNT) as executor:
        futures = [
            executor.submit(material_chunk_from_rows, chunk, headers, code_header, name_header, spec_header, unit_header, analysis_cache)
            for chunk in chunks
        ]
        for future in futures:
            materials.extend(future.result())
    return materials


def spec_quality(value: str) -> int:
    text = clean_cell(value)
    normalized = normalize_text(text)
    if not text:
        return 0
    score = min(len(text), 40)
    if detect_electrical_kind(text):
        score += 35
    if extract_value_key(text, detect_electrical_kind(text)):
        score += 35
    if extract_package(text):
        score += 18
    if extract_tolerance(text):
        score += 14
    if "/" in text:
        score += 8
    if re.fullmatch(r"\d{1,4}/?", normalized):
        score -= 45
    return score


def choose_best_material_spec(row: Dict[str, str], headers: List[str], current_spec: str) -> str:
    best_spec = clean_cell(current_spec)
    best_score = spec_quality(best_spec)
    for header in headers:
        if score_header(header, [], CODE_HEADER_EXCLUDES) < 0:
            continue
        value = clean_cell(row.get(header, ""))
        if not value or value == best_spec:
            continue
        score = spec_quality(value)
        if score > best_score:
            best_spec = value
            best_score = score
    return best_spec


def parse_bom(headers: List[str], rows: List[Dict[str, str]]) -> List[BomItem]:
    spec_header = alias_match(headers, BOM_ALIASES["spec"])
    qty_header = alias_match(headers, BOM_ALIASES["qty"])
    designator_header = alias_match(headers, BOM_ALIASES["designator"])
    footprint_header = alias_match(headers, BOM_ALIASES["footprint"])

    bom_items: List[BomItem] = []
    for row in rows:
        spec = row.get(spec_header, "") if spec_header else first_existing(row, headers, BOM_ALIASES["spec"])
        qty = row.get(qty_header, "") if qty_header else infer_quantity(row, headers)
        designator = row.get(designator_header, "") if designator_header else infer_designator(row, headers)
        footprint = row.get(footprint_header, "") if footprint_header else ""
        if any(row.values()):
            bom_items.append(BomItem(original=row, spec=spec, qty=qty or "1", designator=designator, footprint=footprint))
    return bom_items


def infer_quantity(row: Dict[str, str], headers: List[str]) -> str:
    for value in row.values():
        if re.fullmatch(r"\d+", value or ""):
            return value
    designator = infer_designator(row, headers)
    if designator:
        refs = [part for part in re.split(r"[,，;\s]+", designator) if part]
        if refs:
            return str(len(refs))
    return "1"


def infer_designator(row: Dict[str, str], headers: List[str]) -> str:
    for header in headers:
        value = row.get(header, "")
        if re.search(r"\b[A-Z]{1,4}\d+\b", value or "", flags=re.I):
            return value
    return ""


def detect_process(bom: BomItem) -> str:
    text = normalize_text(" ".join([bom.spec, bom.footprint, bom.designator]))
    for keyword in POST_WELD_KEYWORDS:
        if normalize_text(keyword) in text:
            return "2"
    return "1"


def value_number_key(value: float) -> str:
    return f"{value:.12g}"


def normalize_resistance_value(number: str, unit: str) -> Optional[str]:
    try:
        value = float(number)
    except ValueError:
        return None
    unit = unit.lower()
    if unit == "r":
        multiplier = 1
    elif unit == "k":
        multiplier = 1000
    elif unit == "m":
        multiplier = 1000000
    else:
        return None
    return value_number_key(value * multiplier)


def normalize_si_value(number: str, unit: str, kind: str) -> Optional[str]:
    try:
        value = float(number)
    except ValueError:
        return None
    unit = unit.lower()
    multipliers = {
        "capacitor": {"pf": 1e-12, "nf": 1e-9, "uf": 1e-6, "f": 1},
        "inductor": {"nh": 1e-9, "uh": 1e-6, "mh": 1e-3, "h": 1},
    }
    multiplier = multipliers.get(kind, {}).get(unit)
    if multiplier is None:
        return None
    return value_number_key(value * multiplier)


def normalize_eia_value(code: str, kind: str) -> Optional[str]:
    """Normalize isolated three/four digit EIA markings by component type."""
    if kind == "capacitor" and len(code) == 3:
        significant_digits = code[:2]
        exponent = int(code[2])
        if significant_digits == "00":
            return None
        value_pf = int(significant_digits) * (10**exponent)
        return normalize_si_value(str(value_pf), "pf", "capacitor")
    if kind == "resistor" and len(code) in (3, 4):
        significant_digits = code[:-1]
        exponent = int(code[-1])
        if int(significant_digits) == 0:
            return value_number_key(0)
        return value_number_key(int(significant_digits) * (10**exponent))
    return None


def extract_isolated_eia_code(text: str, kind: str) -> Optional[str]:
    raw_text = str(text or "").lower()
    lengths = "3" if kind == "capacitor" else "3,4"
    for match in re.finditer(rf"(?<![a-z0-9])(\d{{{lengths}}})(?![a-z0-9])", raw_text):
        normalized = normalize_eia_value(match.group(1), kind)
        if normalized is not None:
            return normalized
    return None


def extract_value_key(text: str, kind: Optional[str]) -> Optional[str]:
    normalized = normalize_text(text)
    if kind in {"resistor", "thermistor", "potentiometer"}:
        raw_resistance = str(text or "").lower().replace("ω", "ohm").replace("Ω", "ohm")
        milliohm = re.search(r"(?<![a-z0-9.])(\d+(?:\.\d+)?)\s*m(?:r|ohm)(?![a-z0-9])", raw_resistance)
        if milliohm:
            return value_number_key(float(milliohm.group(1)) / 1000.0)
        match = re.search(r"\b(\d+(?:\.\d+)?)\s*([rkm])(?=\b|[%/])", normalized)
        if match:
            return normalize_resistance_value(match.group(1), match.group(2))
        raw_text = str(text or "").lower().replace("ω", "r").replace("Ω", "r")
        embedded = re.search(r"(?<![a-z0-9])(\d*)([rkm])(\d+)(?![a-z0-9])", raw_text)
        if embedded:
            left, unit, right = embedded.groups()
            if not left and unit == "r" and right in CHIP_PACKAGES:
                embedded = None
            else:
                return normalize_resistance_value(f"{left or '0'}.{right}", unit)
        eia_value = extract_isolated_eia_code(text, "resistor")
        if eia_value is not None:
            return eia_value
    if kind == "capacitor":
        match = re.search(r"\b(\d+(?:\.\d+)?)\s*(pf|nf|uf|f)(?=\b|/)", normalized)
        if match:
            return normalize_si_value(match.group(1), match.group(2), "capacitor")
        eia_value = extract_isolated_eia_code(text, "capacitor")
        if eia_value is not None:
            return eia_value
    if kind == "inductor":
        match = re.search(r"\b(\d+(?:\.\d+)?)\s*(uh|mh|nh|h)(?=\b|/)", normalized)
        if match:
            return normalize_si_value(match.group(1), match.group(2), "inductor")
    if kind == "crystal":
        match = re.search(r"(?<![a-z0-9.])(\d+(?:\.\d+)?)\s*(ghz|mhz|khz|hz)(?![a-z])", normalized)
        if match:
            multiplier = {"hz": 1.0, "khz": 1e3, "mhz": 1e6, "ghz": 1e9}[match.group(2)]
            return value_number_key(float(match.group(1)) * multiplier)
    return None


def _canonical_dimensions(width: float, height: float) -> Tuple[float, float]:
    """Orientation does not change the body-size identity of a package."""
    return (width, height) if width <= height else (height, width)


def _package_key(family: Optional[str], pins: Optional[int], width: Optional[float], height: Optional[float]) -> Optional[str]:
    if not family:
        return None
    result = family
    if pins:
        result += str(pins)
    if width is not None and height is not None:
        width, height = _canonical_dimensions(width, height)
        result += f"@{value_number_key(width)}x{value_number_key(height)}"
    return result


def extract_package_details(text: str) -> Tuple[Optional[str], Optional[str], Optional[int], Optional[float], Optional[float]]:
    """Parse package aliases into canonical family, pin count and body size.

    Examples that intentionally collapse to one key include QFN32-4*4,
    QFN4X4 32 and QFN32-4MM.  Details remain separate so incomplete package
    descriptions can be reviewed instead of being treated as a false conflict.
    """
    raw = str(text or "").lower()
    raw = raw.replace("μ", "u").replace("µ", "u").replace("×", "x").replace("＊", "x").replace("*", "x")
    raw = re.sub(r"(?<=\d)x+(?=\d)", "x", raw)
    raw = raw.replace("–", "-").replace("—", "-").replace("－", "-")
    separated = re.sub(r"[-_/,:;()\[\]{}]+", " ", raw)
    separated = re.sub(r"\s+", " ", separated)

    family_match = re.search(rf"(?<![a-z])({PACKAGE_FAMILY_PATTERN})(?![a-z])", separated)
    if family_match:
        family = PACKAGE_FAMILY_ALIASES[family_match.group(1)]
        tail = separated[family_match.end(): family_match.end() + 42]
        pins: Optional[int] = None
        width: Optional[float] = None
        height: Optional[float] = None

        # Chinese footprint libraries also compact 3.3x3.3 into 3333:
        # PDFN3333-8 and QFN4040-32.
        compact_dimensions = re.match(r"\s*(\d{2})(\d{2})\s+(\d{1,3})(?:\s*(?:pin|pins|p|lead|leads))?", tail)
        if compact_dimensions:
            width = int(compact_dimensions.group(1)) / 10.0
            height = int(compact_dimensions.group(2)) / 10.0
            pins = int(compact_dimensions.group(3))

        # QFN4x4 32 / DFN3x3-8: dimensions before lead count.
        match = None if compact_dimensions else re.match(r"\s*(\d+(?:\.\d+)?)\s*x\s*(\d+(?:\.\d+)?)(?:\s*mm)?(?:\s+(\d{1,3})(?:\s*(?:pin|pins|p|lead|leads))?)?", tail)
        if match:
            width, height = float(match.group(1)), float(match.group(2))
            pins = int(match.group(3)) if match.group(3) else None
        elif not compact_dimensions:
            # QFN32-4x4 / QFN-32-1EP_4x4mm.
            match = re.match(r"\s*(\d{1,3})(?:\s*(?:pin|pins|p|lead|leads))?(?:\s+\d+\s*ep)?\s+(\d+(?:\.\d+)?)\s*x\s*(\d+(?:\.\d+)?)(?:\s*mm)?", tail)
            if match:
                pins = int(match.group(1))
                width, height = float(match.group(2)), float(match.group(3))
            else:
                # A widely used square-body shorthand: QFN32-4MM.
                match = re.match(r"\s*(\d{1,3})(?:\s*(?:pin|pins|p|lead|leads))?\s+(\d+(?:\.\d+)?)\s*mm(?!\s*x)", tail)
                if match:
                    pins = int(match.group(1))
                    width = height = float(match.group(2))
                else:
                    match = re.match(r"\s*(\d{1,3})(?:\s*(?:pin|pins|p|lead|leads))?", tail)
                    if match:
                        pins = int(match.group(1))

        # QFN37-5555 is a 37-lead, 5.5 x 5.5 mm shorthand.
        pin_compact_dimensions = re.match(r"\s*(\d{1,3})\s+(\d{2})(\d{2})(?!\d)", tail)
        if pin_compact_dimensions and int(pin_compact_dimensions.group(1)) <= 512:
            pins = int(pin_compact_dimensions.group(1))
            width = int(pin_compact_dimensions.group(2)) / 10.0
            height = int(pin_compact_dimensions.group(3)) / 10.0

        # EasyEDA/JLCPCB exports dimensions as QFN-28_L5.0-W5.0-P0.50.
        lcsc_dimensions = re.search(r"(?:^|\s)l\s*(\d+(?:\.\d+)?)\s+w\s*(\d+(?:\.\d+)?)", tail)
        if lcsc_dimensions:
            width, height = float(lcsc_dimensions.group(1)), float(lcsc_dimensions.group(2))

        if width is not None and height is not None:
            if not (0.4 <= width <= 100.0 and 0.4 <= height <= 100.0):
                width = height = None
            else:
                width, height = _canonical_dimensions(width, height)
        return _package_key(family, pins, width, height), family, pins, width, height

    # Transistor and diode package names have a numeric family code and may
    # append the actual terminal count (SOT23-6, SOD-123, TO-252/DPAK).
    special = re.search(r"(?<![a-z0-9])(tsot|sot|sod|sc|to)\s*(\d{2,3})(?:\s+(\d{1,2})(?:\s*(?:pin|pins|p))?)?(?!\d)", separated)
    if special:
        family = f"{special.group(1)}{special.group(2)}"
        pins = int(special.group(3)) if special.group(3) else None
        package_aliases = {
            "tsot23": ("sot23", None), "sot25": ("sot23", 5), "sot26": ("sot23", 6),
            "sot323": ("sc70", 3), "sot353": ("sc70", 5), "sot363": ("sc70", 6),
            "sot666": ("sot666", 6), "sod123": ("sod123", 2),
        }
        if family in package_aliases:
            family, default_pins = package_aliases[family]
            pins = pins or default_pins
        tail = separated[special.end(): special.end() + 42]
        dimensions = re.search(r"(?:^|\s)l\s*(\d+(?:\.\d+)?)\s+w\s*(\d+(?:\.\d+)?)", tail)
        width = height = None
        if dimensions:
            width, height = _canonical_dimensions(float(dimensions.group(1)), float(dimensions.group(2)))
        return _package_key(family, pins, width, height), family, pins, width, height
    power_aliases = {
        "dpak": "to252", "d-pak": "to252", "d2pak": "to263", "d²pak": "to263",
        "sma": "sma", "smb": "smb", "smc": "smc", "melf": "melf", "minimelf": "minimelf",
    }
    for alias, family in power_aliases.items():
        if re.search(rf"(?<![a-z0-9]){re.escape(alias)}(?![a-z0-9])", raw):
            return family, family, None, None, None

    # Imperial passive packages may carry a type prefix/suffix: C0603,
    # R0603, F1206, 0603-LED, SMT0603 and 0603SMD.
    if re.search(r"(?<!\d)0603\s*m(?![a-z0-9])", raw):
        return "0201", "0201", None, None, None
    for metric_package, imperial_package in METRIC_PACKAGE_ALIASES.items():
        escaped = re.escape(metric_package)
        metric_pattern = rf"(?:smt|smd|[crlf])[-_ ]*{escaped}(?!\d)|(?<!\d){escaped}[-_ ]*(?:smt|smd)|(?:封装|package)\s*[:：]?\s*{escaped}(?!\d)"
        if re.search(metric_pattern, raw):
            return imperial_package, imperial_package, None, None, None
    passive = re.search(r"(?<!\d)(0201|0402|0603|0805|1206|1210|1812|2010|2512)(?!\d)", raw)
    if passive:
        package = passive.group(1)
        return package, package, None, None, None

    # Radial capacitors, inductors and custom modules commonly expose only
    # body dimensions (C6.3x8, L21x12-CB, 5x7 mm).
    dimension = re.search(r"(?<!\d)(\d+(?:\.\d+)?)\s*x\s*(\d+(?:\.\d+)?)(?:\s*mm)?(?!\d)", raw)
    if dimension:
        width, height = _canonical_dimensions(float(dimension.group(1)), float(dimension.group(2)))
        if 0.4 <= width <= 100.0 and 0.4 <= height <= 100.0:
            family = "body"
            return _package_key(family, None, width, height), family, None, width, height
    return None, None, None, None, None


def extract_package(text: str) -> Optional[str]:
    return extract_package_details(text)[0]


def extract_package_relaxed(text: str) -> Optional[str]:
    direct = extract_package(text)
    if direct:
        return direct
    compact = normalize_key(normalize_text(text))
    for package in CHIP_PACKAGES:
        if re.search(rf"(?<!\d){re.escape(package)}(?!\d)", compact):
            return package
    return None


def extract_pin_count(text: str, kind: Optional[str], package_pins: Optional[int] = None) -> Optional[int]:
    if package_pins:
        return package_pins
    raw = str(text or "").lower().replace("－", "-")
    match = re.search(r"(?<!\d)(\d{1,3})\s*(?:pin|pins|p|位)(?![a-z])", raw)
    if not match:
        match = re.search(r"(?:pin|pins|引脚)\s*[:：-]?\s*(\d{1,3})(?!\d)", raw)
    if match:
        pins = int(match.group(1))
        return pins if 1 <= pins <= 512 else None
    if kind == "connector":
        match = re.search(r"(?<![a-z0-9])(?:con|conn|cn)\s*[-_]?\s*(\d{1,3})(?![a-z0-9])", raw)
        if match:
            return int(match.group(1))
        match = re.search(r"(?<!\d)([12])\s*x\s*0?(\d{1,2})(?!\d)", raw)
        if match:
            return int(match.group(1)) * int(match.group(2))
    return None


def extract_current(text: str, kind: Optional[str]) -> Optional[float]:
    if kind != "fuse":
        return None
    normalized = normalize_text(text)
    match = re.search(r"(?<![a-z0-9.])(\d+(?:\.\d+)?)\s*(ma|a|amp|amps)(?![a-z])", normalized)
    if not match:
        return None
    current = float(match.group(1))
    return current / 1000.0 if match.group(2) == "ma" else current


def extract_component_subtype(text: str, kind: Optional[str]) -> Optional[str]:
    raw = str(text or "").lower().replace("_", " ")
    compact = normalize_key(raw)
    if kind == "connector":
        connector_aliases = (
            ("usb-type-c", ("typec", "usbc", "usbtypec")),
            ("micro-usb", ("microusb",)),
            ("mini-usb", ("miniusb",)),
            ("usb-a", ("usba",)),
            ("jst", ("jst",)),
            ("header", ("pinheader", "header", "排针", "排母")),
            ("terminal", ("terminal", "端子")),
        )
        for canonical, aliases in connector_aliases:
            if any(alias in compact or alias in raw for alias in aliases):
                return canonical
    if kind == "led":
        colors = {
            "green": ("green", "绿色", "绿灯"), "red": ("red", "红色", "红灯"),
            "blue": ("blue", "蓝色", "蓝灯"), "yellow": ("yellow", "黄色", "黄灯"),
            "white": ("white", "白色", "白灯"), "amber": ("amber", "橙色", "琥珀"),
            "rgb": ("rgb", "全彩", "三色"),
        }
        for canonical, aliases in colors.items():
            if any(alias in raw for alias in aliases):
                return canonical
    if kind in {"mosfet", "transistor"}:
        if any(token in compact for token in ("nmos", "nchannel", "mosn", "npn")):
            return "n"
        if any(token in compact for token in ("pmos", "pchannel", "mosp", "pnp")):
            return "p"
    if kind == "diode":
        if "schottky" in raw or "肖特基" in raw:
            return "schottky"
        if "zener" in raw or "稳压" in raw:
            return "zener"
    return None


def extract_tolerance(text: str) -> Optional[str]:
    normalized = normalize_text(text)
    match = re.search(r"(?<![\d.])(\d+(?:\.\d+)?)\s*%", normalized)
    if match:
        return value_number_key(float(match.group(1))) + "%"
    tolerance_words = {
        "d": "0.5%",
        "j": "5%",
        "f": "1%",
        "g": "2%",
        "k": "10%",
        "m": "20%",
    }
    raw_text = str(text or "").lower()
    for token, value in tolerance_words.items():
        if re.search(rf"(?<![a-z0-9]){token}(?![a-z0-9])", raw_text):
            return value
    return None


def tolerance_percent_from_key(tolerance: Optional[str]) -> Optional[float]:
    if not tolerance:
        return None
    try:
        return float(tolerance.rstrip("%"))
    except ValueError:
        return None


def extract_voltage(text: str, kind: Optional[str] = None) -> Optional[float]:
    normalized = normalize_text(text)
    match = re.search(r"(?<![\d.])(\d+(?:\.\d+)?)\s*(kv|v)(?:dc)?(?=\b|[/,;+\-)])", normalized)
    if match:
        voltage = float(match.group(1))
        if match.group(2) == "kv":
            voltage *= 1000.0
        return voltage
    match = re.search(r"(?<![\d.])(\d+(?:\.\d+)?)\s*wv(?=\b|[/,;+\-)])", normalized)
    if match:
        return float(match.group(1))
    match = re.search(r"\bwv\s*(\d+(?:\.\d+)?)(?=\b|[/,;+\-)])", normalized)
    if match:
        return float(match.group(1))
    if kind == "capacitor":
        raw_text = str(text or "").lower()
        for code, voltage in CAPACITOR_VOLTAGE_CODES.items():
            if re.search(rf"(?<![a-z0-9]){re.escape(code)}(?![a-z0-9])", raw_text):
                return voltage
    return None


def extract_dielectric(text: str) -> Optional[str]:
    normalized = normalize_key(str(text or "").lower())
    aliases = {
        "np0": "c0g",
        "npo": "c0g",
        "c0g": "c0g",
        "x7r": "x7r",
        "x5r": "x5r",
        "y5v": "y5v",
        "z5u": "z5u",
    }
    for token, canonical in aliases.items():
        if token in normalized:
            return canonical
    return None


def component_signature(text: str) -> ComponentSignature:
    kind = detect_electrical_kind(text)
    tolerance = extract_tolerance(text)
    package, package_family, package_pins, package_width, package_height = extract_package_details(text)
    return ComponentSignature(
        kind=kind,
        value_key=extract_value_key(text, kind),
        package=package or extract_package_relaxed(text),
        package_family=package_family,
        pin_count=extract_pin_count(text, kind, package_pins),
        package_width_mm=package_width,
        package_height_mm=package_height,
        tolerance=tolerance,
        tolerance_percent=tolerance_percent_from_key(tolerance),
        voltage_v=extract_voltage(text, kind),
        current_a=extract_current(text, kind),
        dielectric=extract_dielectric(text),
        subtype=extract_component_subtype(text, kind),
    )


def material_component_signature(material: MaterialItem) -> ComponentSignature:
    semantic_values = [material.spec, material.name]
    semantic_values.extend(
        clean_cell(value)
        for value in material.source.values()
        if clean_cell(value) and clean_cell(value) != clean_cell(material.code)
    )
    semantic_text = " ".join(semantic_values)
    package_text = " ".join([material.spec, material.code, semantic_text])
    kind_text = " ".join([material.name, material.spec, material.code])
    kind = detect_electrical_kind(kind_text)
    tolerance = extract_tolerance(semantic_text)
    package, package_family, package_pins, package_width, package_height = extract_package_details(package_text)
    return ComponentSignature(
        kind=kind,
        value_key=extract_value_key(material.spec or semantic_text, kind),
        package=package or extract_package_relaxed(package_text),
        package_family=package_family,
        pin_count=extract_pin_count(package_text, kind, package_pins),
        package_width_mm=package_width,
        package_height_mm=package_height,
        tolerance=tolerance,
        tolerance_percent=tolerance_percent_from_key(tolerance),
        voltage_v=extract_voltage(semantic_text, kind),
        current_a=extract_current(semantic_text, kind),
        dielectric=extract_dielectric(semantic_text),
        subtype=extract_component_subtype(semantic_text, kind),
    )


def apply_material_analysis(
    material: MaterialItem,
    analysis_cache: Optional[Dict[str, Tuple[ComponentSignature, frozenset[str]]]] = None,
) -> None:
    material.normalized_searchable = normalize_text(material.searchable)
    cache_key = material_analysis_key(material)
    cached = analysis_cache.get(cache_key) if analysis_cache else None
    if cached is not None:
        material.signature, material.identifiers = cached
        material.analysis_cache_hit = True
        return
    material.signature = material_component_signature(material)
    material.identifiers = frozenset(extract_part_identifiers(material.searchable))


def bom_signature(bom: BomItem) -> ComponentSignature:
    base_text = " ".join([bom.spec, bom.footprint, bom.designator])
    kind = detect_electrical_kind(base_text)
    designator_match = re.search(r"(?:^|[\s,;])([a-z]{1,4})\s*\d+", str(bom.designator or "").lower())
    if designator_match:
        designator_kind = {
            "f": "fuse", "fu": "fuse", "led": "led", "j": "connector",
            "jp": "connector", "p": "connector", "cn": "connector", "con": "connector", "usb": "connector",
            "u": "ic", "ic": "ic", "q": "transistor", "d": "diode",
            "y": "crystal", "x": "crystal", "sw": "switch", "key": "switch", "k": "relay",
            "r": "resistor", "c": "capacitor", "ec": "capacitor", "l": "inductor",
            "ae": "antenna", "ant": "antenna", "tp": "testpoint",
            "lcm": "module", "lcd": "module", "disp": "module",
        }.get(designator_match.group(1))
        if designator_kind == "fuse":
            has_current_rating = bool(re.search(r"(?<![a-z0-9.])\d+(?:\.\d+)?\s*(?:ma|a)(?![a-z])", str(bom.spec or "").lower()))
            if has_current_rating or kind != "resistor":
                kind = "fuse"
        elif designator_kind in {"led", "connector", "switch", "antenna", "testpoint"}:
            kind = designator_kind
        elif designator_kind == "ic" and kind not in {"connector", "module"}:
            kind = "ic"
        elif designator_kind == "transistor" and kind not in {"mosfet"}:
            kind = "transistor"
        elif designator_kind == "diode" and kind not in {"tvs", "led"}:
            kind = "diode"
        elif kind is None and designator_kind:
            kind = designator_kind
    tolerance = extract_tolerance(bom.spec)
    voltage_v = extract_voltage(bom.spec, kind)
    if kind == "capacitor" and voltage_v is None:
        voltage_v = 50.0
    if kind == "resistor" and tolerance is None:
        tolerance = "1%"
    package, package_family, package_pins, package_width, package_height = extract_package_details(base_text)
    return ComponentSignature(
        kind=kind,
        value_key=extract_value_key(bom.spec, kind),
        package=package or extract_package_relaxed(base_text),
        package_family=package_family,
        pin_count=extract_pin_count(base_text, kind, package_pins),
        package_width_mm=package_width,
        package_height_mm=package_height,
        tolerance=tolerance,
        tolerance_percent=tolerance_percent_from_key(tolerance),
        voltage_v=voltage_v,
        current_a=extract_current(bom.spec, kind),
        dielectric=extract_dielectric(bom.spec),
        subtype=extract_component_subtype(base_text, kind),
    )


def package_dimensions_equal(left: ComponentSignature, right: ComponentSignature) -> bool:
    if None in (left.package_width_mm, left.package_height_mm, right.package_width_mm, right.package_height_mm):
        return False
    left_size = _canonical_dimensions(float(left.package_width_mm), float(left.package_height_mm))
    right_size = _canonical_dimensions(float(right.package_width_mm), float(right.package_height_mm))
    return abs(left_size[0] - right_size[0]) <= 0.12 and abs(left_size[1] - right_size[1]) <= 0.12


def package_has_dimensions(signature: ComponentSignature) -> bool:
    return signature.package_width_mm is not None and signature.package_height_mm is not None


def kinds_compatible(left: Optional[str], right: Optional[str]) -> bool:
    if not left or not right or left == right:
        return True
    compatible_groups = (
        {"transistor", "mosfet"},
        {"diode", "tvs", "led"},
        {"resistor", "thermistor", "varistor", "potentiometer"},
    )
    return any(left in group and right in group for group in compatible_groups)


def signature_compatibility(query_sig: ComponentSignature, candidate_sig: ComponentSignature) -> Tuple[bool, bool, int]:
    """Return hard-conflict, review-warning and maximum allowed score.

    Hard conflicts must never be rescued by fuzzy text similarity. Missing
    critical information is capped below the auto-match threshold so it stays
    available for manual search without being applied automatically.
    """
    if query_sig.kind and candidate_sig.kind and not kinds_compatible(query_sig.kind, candidate_sig.kind):
        return True, False, 0
    if query_sig.value_key and candidate_sig.value_key and query_sig.value_key != candidate_sig.value_key:
        return True, False, 0
    if query_sig.package_family and candidate_sig.package_family:
        if query_sig.package_family != candidate_sig.package_family:
            return True, False, 0
        if query_sig.pin_count and candidate_sig.pin_count and query_sig.pin_count != candidate_sig.pin_count:
            return True, False, 0
        if package_has_dimensions(query_sig) and package_has_dimensions(candidate_sig) and not package_dimensions_equal(query_sig, candidate_sig):
            return True, False, 0
    elif query_sig.package and candidate_sig.package and query_sig.package != candidate_sig.package:
        return True, False, 0
    if query_sig.voltage_v is not None and candidate_sig.voltage_v is not None:
        if candidate_sig.voltage_v + 1e-9 < query_sig.voltage_v:
            return True, False, 0
    if query_sig.current_a is not None and candidate_sig.current_a is not None:
        if abs(query_sig.current_a - candidate_sig.current_a) > max(0.02, query_sig.current_a * 0.03):
            return True, False, 0
    if query_sig.tolerance_percent is not None and candidate_sig.tolerance_percent is not None:
        if candidate_sig.tolerance_percent > query_sig.tolerance_percent + 1e-9:
            return True, False, 0
    if query_sig.dielectric and candidate_sig.dielectric and query_sig.dielectric != candidate_sig.dielectric:
        return True, False, 0

    warning = False
    score_cap = 100
    if query_sig.kind and not candidate_sig.kind:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 8)
    if query_sig.value_key and not candidate_sig.value_key:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 17)
    if query_sig.package and not candidate_sig.package:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 8)
    if query_sig.pin_count and not candidate_sig.pin_count:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 5)
    if package_has_dimensions(query_sig) and not package_has_dimensions(candidate_sig):
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 3)
    if query_sig.voltage_v is not None:
        if candidate_sig.voltage_v is None:
            warning = True
            score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 4)
        elif candidate_sig.voltage_v > query_sig.voltage_v + 1e-9:
            # A higher voltage rating can be usable when the package is equal,
            # but is intentionally highlighted for a human check.
            warning = True
    if query_sig.current_a is not None and candidate_sig.current_a is None:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 8)
    if query_sig.tolerance_percent is not None and candidate_sig.tolerance_percent is None:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 2)
    if query_sig.dielectric and not candidate_sig.dielectric:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 2)
    if query_sig.subtype and candidate_sig.subtype and query_sig.subtype != candidate_sig.subtype:
        return True, False, 0
    if query_sig.subtype and not candidate_sig.subtype:
        warning = True
        score_cap = min(score_cap, AUTO_MATCH_THRESHOLD - 4)
    return False, warning, score_cap


def signatures_have_warning(query_sig: ComponentSignature, candidate_sig: ComponentSignature) -> bool:
    _conflict, warning, _score_cap = signature_compatibility(query_sig, candidate_sig)
    return warning


def extract_part_identifiers(text: str) -> set[str]:
    """Extract model/part-number tokens used as generic-component guards."""
    identifiers: set[str] = set()
    normalized = str(text or "").lower().replace("μ", "u").replace("µ", "u").replace("×", "x").replace("*", "x")
    normalized = re.sub(r"(?<=\d)x+(?=\d)", "x", normalized)
    for raw_token in re.findall(r"[a-z0-9][a-z0-9._-]{2,}", normalized):
        token = normalize_key(raw_token)
        if len(token) < 4 or not re.search(r"[a-z]", token) or not re.search(r"\d", token):
            continue
        if re.fullmatch(r"\d+(?:\d|p|\.)*(?:pf|nf|uf|nh|uh|mh|v|kv|r|k|m)", token):
            continue
        if re.fullmatch(r"[crlf]?(?:0201|0402|0603|0805|1206|1210|1812|2010|2512)", token):
            continue
        if re.fullmatch(r"(?:sot|sod|sc|to)\d{2,3}(?:\d+x\d+)?(?:mm)?", token):
            continue
        if re.fullmatch(r"(?:sop|soic|tssop|ssop|msop|qfn|vqfn|wqfn|dfn|wdfn|son|lga|lqfp|tqfp|qfp|dip|bga|csp|plcc)(?:\d+x\d+|\d+)?(?:mm)?", token):
            continue
        if re.fullmatch(r"(?:[crlf]|smt|smd)?(?:0201|0402|0603|0805|1206|1210|1812|2010|2512)(?:led|smd|smt|[crlf])?", token):
            continue
        if re.fullmatch(r"\d+(?:pin|pins)(?:pad|smd|smt)?", token):
            continue
        if re.fullmatch(r"(?:led|fuse|res|cap|ind)\d+", token):
            continue
        if re.search(r"(?:0201|0402|0603|0805|1206|1210|1812|2010|2512)", token) and any(
            word in token for word in ("led", "smd", "smt", "green", "red", "blue", "white", "yellow")
        ):
            continue
        if re.fullmatch(r"x\d[rs]", token):
            continue
        identifiers.add(token)
    return identifiers


def identifier_family_matches(query_identifiers: set[str], candidate_identifiers: set[str]) -> set[Tuple[str, str]]:
    """Return conservative base-model matches such as SC8815 -> SC8815QDER.

    Six characters plus an equal physical package are required by the caller;
    short connector labels and generic package strings never use this path.
    """
    matches: set[Tuple[str, str]] = set()
    for query_identifier in query_identifiers:
        for candidate_identifier in candidate_identifiers:
            shorter, longer = sorted((query_identifier, candidate_identifier), key=len)
            if len(shorter) >= 6 and len(longer) - len(shorter) <= 10 and longer.startswith(shorter):
                matches.add((query_identifier, candidate_identifier))
    return matches


def structured_match_score(query: str, candidate: str, query_sig: Optional[ComponentSignature] = None, candidate_sig: Optional[ComponentSignature] = None) -> Tuple[int, bool]:
    query_sig = query_sig or component_signature(query)
    candidate_sig = candidate_sig or component_signature(candidate)
    hard_conflict, warning, score_cap = signature_compatibility(query_sig, candidate_sig)
    if hard_conflict:
        return 0, False

    score = 0
    if query_sig.kind and candidate_sig.kind:
        if query_sig.kind == candidate_sig.kind:
            score += 15
        elif kinds_compatible(query_sig.kind, candidate_sig.kind):
            score += 8
    if query_sig.value_key and candidate_sig.value_key and query_sig.value_key == candidate_sig.value_key:
        score += 35
    if query_sig.package and candidate_sig.package and query_sig.package == candidate_sig.package:
        score += 20
    elif query_sig.package_family and candidate_sig.package_family and query_sig.package_family == candidate_sig.package_family:
        score += 8
        if query_sig.pin_count and candidate_sig.pin_count and query_sig.pin_count == candidate_sig.pin_count:
            score += 7
        if package_dimensions_equal(query_sig, candidate_sig):
            score += 5
    if query_sig.pin_count and candidate_sig.pin_count and query_sig.pin_count == candidate_sig.pin_count:
        if not (query_sig.package and candidate_sig.package and query_sig.package == candidate_sig.package):
            score += 15
    if query_sig.voltage_v is not None and candidate_sig.voltage_v is not None:
        if abs(query_sig.voltage_v - candidate_sig.voltage_v) <= 1e-9:
            score += 12
        else:
            score += 6
    if query_sig.current_a is not None and candidate_sig.current_a is not None:
        if abs(query_sig.current_a - candidate_sig.current_a) <= max(0.02, query_sig.current_a * 0.03):
            score += 37
    if query_sig.tolerance_percent is not None and candidate_sig.tolerance_percent is not None:
        if abs(query_sig.tolerance_percent - candidate_sig.tolerance_percent) <= 1e-9:
            score += 7
        elif candidate_sig.tolerance_percent < query_sig.tolerance_percent:
            score += 4
    if query_sig.dielectric and candidate_sig.dielectric and query_sig.dielectric == candidate_sig.dielectric:
        score += 7
    if query_sig.subtype and candidate_sig.subtype and query_sig.subtype == candidate_sig.subtype:
        score += 45

    normalized_query = normalize_text(query)
    normalized_candidate = normalize_text(candidate)
    if normalized_query and normalized_candidate:
        score += int(10 * SequenceMatcher(None, normalized_query, normalized_candidate).ratio())

    return max(0, min(100, score, score_cap)), warning


def match_score(bom_text: str, material: MaterialItem) -> int:
    score, _warning = match_score_detail(bom_text, material)
    return score


def match_score_detail(bom_text: str, material: MaterialItem, query_sig: Optional[ComponentSignature] = None) -> Tuple[int, bool]:
    query = normalize_text(bom_text)
    spec = normalize_text(material.spec)
    combined = material.normalized_searchable or normalize_text(material.searchable)
    if not query or not combined:
        return 0, False
    query_sig = query_sig or component_signature(bom_text)
    candidate_sig = material.signature or material_component_signature(material)
    compatibility = 1.0 if kinds_compatible(query_sig.kind, candidate_sig.kind) else 0.0
    if compatibility == 0:
        return 0, False
    hard_conflict, warning, score_cap = signature_compatibility(query_sig, candidate_sig)
    if hard_conflict:
        return 0, False
    structured_score, warning = structured_match_score(bom_text, material.searchable, query_sig, candidate_sig)
    identity_guard_kinds = {"ic", "mosfet", "transistor", "diode", "tvs"}
    query_identifiers = extract_part_identifiers(bom_text)
    candidate_identifiers = material.identifiers or extract_part_identifiers(material.searchable)
    matching_identifiers = query_identifiers & candidate_identifiers
    family_identifiers = identifier_family_matches(query_identifiers, set(candidate_identifiers))
    package_is_equal = bool(query_sig.package and candidate_sig.package and query_sig.package == candidate_sig.package)
    family_identity_is_usable = bool(family_identifiers and package_is_equal)
    if query_sig.kind in identity_guard_kinds and query_identifiers and not matching_identifiers and not family_identity_is_usable:
        return 0, False
    if matching_identifiers:
        structured_score = min(score_cap, structured_score + 55)
    elif family_identity_is_usable:
        structured_score = min(score_cap, structured_score + 42)
        warning = True
    if any(
        value is not None
        for value in (
            query_sig.kind,
            query_sig.value_key,
            query_sig.package,
            query_sig.pin_count,
            query_sig.tolerance_percent,
            query_sig.voltage_v,
            query_sig.current_a,
            query_sig.dielectric,
            query_sig.subtype,
        )
    ):
        # Structured electrical/physical fields own the decision. Text
        # similarity is already a small tie-breaker inside the score and can no
        # longer override a package, value, voltage or tolerance conflict.
        return structured_score, warning

    # For ICs, diodes, connectors and other generic components, an explicit
    # model token is a hard identity guard. For example SMAJ30CA must not be
    # auto-replaced by the textually similar SMAJ33CA.
    query_identifiers = extract_part_identifiers(bom_text)
    if query_identifiers:
        candidate_identifiers = material.identifiers or extract_part_identifiers(material.searchable)
        if not (query_identifiers & candidate_identifiers):
            return 0, False

    if query == spec or query == combined:
        return min(100, score_cap), warning
    if query and spec and (query in spec or spec in query):
        shorter = min(len(query), len(spec))
        longer = max(len(query), len(spec))
        text_score = int(min(99, 86 + int(13 * shorter / max(1, longer))) * compatibility)
        return min(max(text_score, structured_score), score_cap), warning

    query_tokens = set(query.split())
    combined_tokens = set(combined.split())
    token_score = 0
    if query_tokens and combined_tokens:
        token_score = int(100 * len(query_tokens & combined_tokens) / len(query_tokens | combined_tokens))

    spec_ratio = int(SequenceMatcher(None, query, spec).ratio() * 100) if spec else 0
    combined_ratio = int(SequenceMatcher(None, query, combined).ratio() * 100)
    text_score = int(max(token_score, spec_ratio, combined_ratio) * compatibility)
    if structured_score >= 60:
        text_score = max(text_score, structured_score)
    else:
        text_score = max(text_score, structured_score)
    return int(min(100, text_score, score_cap)), warning


def detect_electrical_kind(text: str) -> Optional[str]:
    raw_text = str(text or "").lower()
    normalized = normalize_text(text)
    if re.search(r"(?<![a-z0-9])(?:smaj|smbj|smcj|p?esd)\d", raw_text):
        return "tvs"
    # Named non-passive categories take priority over suffixes such as
    # 1206-R. A fuse described as "8A 32V Fuse ... 1206-R" is still a fuse.
    for kind, keywords in COMPONENT_KIND_KEYWORDS:
        if kind not in {"resistor", "capacitor", "inductor"} and any(
            keyword in raw_text or keyword in normalized for keyword in keywords
        ):
            return kind
    if re.search(r"(?<![a-z0-9])mos(?![a-z0-9])", raw_text):
        return "mosfet"
    # An explicit C/R/L footprint prefix is stronger context than a compact
    # marking. For example, 1H on a C0402 capacitor is a 50 V voltage code, not
    # a one-henry inductance value.
    chip_codes = r"(?:0201|0402|0603|0805|1206|1210|1812|2010|2512)"
    if re.search(rf"(?<![a-z0-9])c[-_ ]*{chip_codes}(?!\d)|(?<!\d){chip_codes}[-_ ]*c(?![a-z0-9])", raw_text):
        return "capacitor"
    if re.search(rf"(?<![a-z0-9])r[-_ ]*{chip_codes}(?!\d)|(?<!\d){chip_codes}[-_ ]*r(?![a-z0-9])", raw_text):
        return "resistor"
    if re.search(rf"(?<![a-z0-9])l[-_ ]*{chip_codes}(?!\d)|(?<!\d){chip_codes}[-_ ]*l(?![a-z0-9])", raw_text):
        return "inductor"

    # Explicit descriptions are stronger than a designator. This lets a
    # Q-designated MOSFET in DFN and a U-designated USB controller keep their
    # real categories.
    for kind, keywords in COMPONENT_KIND_KEYWORDS:
        if any(keyword in raw_text or keyword in normalized for keyword in keywords):
            return kind
    if re.search(r"(?<![a-z0-9])mos(?![a-z0-9])", raw_text):
        return "mosfet"
    if re.search(r"(?<![a-z0-9])(?:con|conn|cn)\s*[-_]?\s*\d+\b", raw_text):
        return "connector"
    if re.search(r"(?<![a-z0-9])led\s*[-_]?\s*\d+\b", raw_text):
        return "led"
    if re.search(r"\b\d+(\.\d+)?\s*(uh|mh|nh|h)(\b|/)", normalized):
        return "inductor"
    if re.search(r"\b\d+(\.\d+)?\s*(pf|nf|uf|f)(\b|/)", normalized):
        return "capacitor"
    if re.search(r"\b\d+(\.\d+)?\s*(r|k|m)(\b|[%/])", normalized) or "ohm" in normalized:
        return "resistor"
    if re.search(r"\bcap\b", raw_text):
        return "capacitor"
    if re.search(r"\bres\b", raw_text):
        return "resistor"

    # Internal material codes often retain an unambiguous category segment.
    code_kind_patterns = {
        "fuse": r"(?:^|[._-])fu(?:[._-]|$)",
        "resistor": r"(?:^|[._-])r(?:[._-]|$)",
        "capacitor": r"(?:^|[._-])c(?:[._-]|$)",
        "inductor": r"(?:^|[._-])l(?:[._-]|$)",
        "diode": r"(?:^|[._-])d(?:[._-]|$)",
        "connector": r"(?:^|[._-])(?:con|cn|j)(?:[._-]|$)",
        "ic": r"(?:^|[._-])(?:ic|u)(?:[._-]|$)",
    }
    for kind, pattern in code_kind_patterns.items():
        if re.search(pattern, raw_text):
            return kind

    # Reference designators cover BOMs whose value is just a model number.
    designator_map = {
        "led": "led", "fu": "fuse", "f": "fuse", "fb": "inductor",
        "r": "resistor", "rt": "thermistor", "ntc": "thermistor", "rv": "varistor",
        "c": "capacitor", "l": "inductor", "d": "diode", "q": "transistor",
        "u": "ic", "ic": "ic", "j": "connector", "jp": "connector", "cn": "connector", "con": "connector",
        "p": "connector", "usb": "connector", "y": "crystal", "x": "crystal",
        "sw": "switch", "key": "switch", "s": "switch", "k": "relay", "t": "transformer",
        "bz": "buzzer", "ls": "buzzer", "bt": "battery", "b": "battery",
        "ae": "antenna", "ant": "antenna", "tp": "testpoint", "ec": "capacitor",
        "lcm": "module", "lcd": "module", "disp": "module",
    }
    designator = re.search(r"(?:^|[\s,;])([a-z]{1,4})\s*\d+(?=$|[\s,;])", raw_text)
    if designator and designator.group(1) in designator_map:
        return designator_map[designator.group(1)]

    # A named leaded package is IC context when nothing more specific exists.
    if re.search(rf"(?<![a-z])(?:{PACKAGE_FAMILY_PATTERN})(?:[-_ ]*\d+|\d+)", raw_text):
        return "ic"
    return None


def unit_compatibility(query: str, candidate: str) -> float:
    query_kind = detect_electrical_kind(query)
    candidate_kind = detect_electrical_kind(candidate)
    if query_kind and candidate_kind and not kinds_compatible(query_kind, candidate_kind):
        return 0.0
    return 1.0


def build_material_index(materials: List[MaterialItem]) -> Dict[Tuple[str, str, str], List[int]]:
    index: Dict[Tuple[str, str, str], List[int]] = {}
    for material_index, material in enumerate(materials):
        signature = material.signature or material_component_signature(material)
        keys = []
        if signature.kind and signature.value_key and signature.package:
            keys.append((signature.kind, signature.value_key, signature.package))
        if signature.kind and signature.value_key:
            keys.append((signature.kind, signature.value_key, ""))
        if signature.kind and signature.package:
            keys.append((signature.kind, "", signature.package))
        if signature.kind:
            keys.append((signature.kind, "", ""))
        for key in keys:
            index.setdefault(key, []).append(material_index)
    return index


def candidate_material_indexes(bom_sig: ComponentSignature, material_index: Dict[Tuple[str, str, str], List[int]], material_count: int) -> List[int]:
    keys = []
    if bom_sig.kind and bom_sig.value_key and bom_sig.package:
        keys.append((bom_sig.kind, bom_sig.value_key, bom_sig.package))
    if bom_sig.kind and bom_sig.value_key:
        keys.append((bom_sig.kind, bom_sig.value_key, ""))
    if bom_sig.kind and bom_sig.package:
        keys.append((bom_sig.kind, "", bom_sig.package))
    if bom_sig.kind:
        keys.append((bom_sig.kind, "", ""))
    for key in keys:
        indexes = material_index.get(key, [])
        if indexes:
            return indexes
    return list(range(material_count))


def narrow_by_tolerance(candidate_indexes, bom_sig: ComponentSignature, materials: List[MaterialItem]) -> List[int]:
    if bom_sig.tolerance_percent is None:
        return list(candidate_indexes)
    acceptable = []
    for index in candidate_indexes:
        signature = materials[index].signature
        if signature is None or signature.tolerance_percent is None:
            acceptable.append(index)
        elif signature.tolerance_percent <= bom_sig.tolerance_percent + 1e-9:
            acceptable.append(index)
    return acceptable


def independently_verify_match(
    query_sig: ComponentSignature,
    candidate_sig: ComponentSignature,
    score: int,
    runner_up_score: int,
) -> Tuple[bool, bool]:
    """Second pass that does not use fuzzy text similarity."""
    hard_conflict, warning, _score_cap = signature_compatibility(query_sig, candidate_sig)
    if hard_conflict or score < AUTO_MATCH_THRESHOLD:
        return False, warning
    if query_sig.value_key and not candidate_sig.value_key:
        return False, True
    if query_sig.package and not candidate_sig.package:
        return False, True
    if query_sig.pin_count and not candidate_sig.pin_count:
        return False, True
    if query_sig.voltage_v is not None and candidate_sig.voltage_v is None:
        return False, True
    if query_sig.current_a is not None and candidate_sig.current_a is None:
        return False, True
    if query_sig.subtype and not candidate_sig.subtype:
        return False, True
    if runner_up_score > 0 and score - runner_up_score < AMBIGUITY_MARGIN:
        if query_sig.kind in {"ic", "mosfet", "transistor", "diode", "tvs"}:
            return False, True
        warning = True
    return True, warning


def best_material_match(bom: BomItem, materials: List[MaterialItem], material_index: Optional[Dict[Tuple[str, str, str], List[int]]] = None) -> Tuple[Optional[MaterialItem], int, bool]:
    query_parts = [bom.spec]
    if bom.footprint:
        query_parts.append(bom.footprint)
    query = " ".join(query_parts)
    query_sig = bom_signature(bom)
    candidate_indexes = candidate_material_indexes(query_sig, material_index or {}, len(materials)) if material_index else range(len(materials))
    candidate_indexes = narrow_by_tolerance(candidate_indexes, query_sig, materials)
    ranked: List[Tuple[int, int, bool, MaterialItem]] = []
    for material_index in candidate_indexes:
        material = materials[material_index]
        score, warning = match_score_detail(query, material, query_sig)
        if score > 0:
            ranked.append((score, -material_index, warning, material))
    if not ranked:
        return None, 0, False
    ranked.sort(reverse=True, key=lambda item: (item[0], item[1]))
    best_score_value, _negative_index, best_warning, best_item = ranked[0]
    runner_up_score = ranked[1][0] if len(ranked) > 1 else 0
    candidate_sig = best_item.signature or material_component_signature(best_item)
    verified, verification_warning = independently_verify_match(
        query_sig,
        candidate_sig,
        best_score_value,
        runner_up_score,
    )
    if best_score_value >= AUTO_MATCH_THRESHOLD and not verified:
        best_score_value = AUTO_MATCH_THRESHOLD - 1
    return best_item, best_score_value, best_warning or verification_warning


def build_result_row(bom: BomItem, material: Optional[MaterialItem], score: int = 0) -> Dict[str, str]:
    result = {
        "层级": "1",
        "工序": detect_process(bom),
        "物料编码": material.code if material and score >= AUTO_MATCH_THRESHOLD else "",
        "物料名称": material.name if material and score >= AUTO_MATCH_THRESHOLD else "",
        "规格": material.spec if material and score >= AUTO_MATCH_THRESHOLD else bom.spec,
        "用量": bom.qty or "1",
        "单位": material.unit if material and material.unit else "pcs",
        "位号": bom.designator,
        "点数": "",
    }
    return result


def export_xlsx(path: str, rows: List[Dict[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "匹配结果"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for col_index, header in enumerate(OUTPUT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, record in enumerate(rows, start=2):
        for col_index, header in enumerate(OUTPUT_COLUMNS, start=1):
            sheet.cell(row=row_index, column=col_index, value=record.get(header, ""))

    for col_index, header in enumerate(OUTPUT_COLUMNS, start=1):
        values = [str(sheet.cell(row=row, column=col_index).value or "") for row in range(1, min(sheet.max_row, 100) + 1)]
        width = min(max(len(value) for value in values) + 4, 40)
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    sheet.freeze_panes = "A2"
    workbook.save(path)


class GridTable(ttk.Frame):
    def __init__(
        self,
        parent,
        row_height: int = 34,
        header_height: int = 26,
        max_lines: int = 2,
        visible_rows: int = 8,
        fill_available_width: bool = False,
        on_select=None,
        on_double_click=None,
        on_checkbox=None,
    ):
        super().__init__(parent)
        self.headers: List[str] = []
        self.display_headers: List[str] = []
        self.rows: List[List[str]] = []
        self.row_ids: List[int] = []
        self.column_widths: List[int] = []
        self.manual_widths: Dict[str, int] = {}
        self.confirmed_ids = set()
        self.auto_matched_ids = set()
        self.warning_ids = set()
        self.selected_index: Optional[int] = None
        self.scroll_after_id = None
        self.resizing_col: Optional[int] = None
        self.resize_start_x = 0
        self.resize_start_width = 0
        self.tooltip_window: Optional[tk.Toplevel] = None
        self.tooltip_after = None
        self.tooltip_cell: Optional[Tuple[int, str]] = None
        self.row_height = row_height
        self.header_height = header_height
        self.max_lines = max_lines
        self.fill_available_width = fill_available_width
        self.on_select = on_select
        self.on_double_click = on_double_click
        self.on_checkbox = on_checkbox
        self.font = tkfont.Font(family="Microsoft YaHei UI", size=10)
        self.header_font = tkfont.Font(family="Microsoft YaHei UI", size=10, weight="bold")
        self.grid_color = "#B9C0C8"
        self.header_bg = "#EEF2F7"
        self.selected_bg = "#BFDFFF"
        self.auto_matched_bg = "#44DD8A"
        self.confirmed_bg = "#33AE3B"
        self.warning_bg = "#FF6B6B"
        self.text_color = "#111111"
        self._resize_after = None

        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)
        canvas_height = self.header_height + self.row_height * visible_rows
        self.canvas = tk.Canvas(self, bg="#FFFFFF", height=canvas_height, highlightthickness=1, highlightbackground="#AAB2BD")
        self.v_scroll = ttk.Scrollbar(self, orient=tk.VERTICAL, command=self._on_y_scroll)
        self.h_scroll = ttk.Scrollbar(self, orient=tk.HORIZONTAL, command=self._on_x_scroll)
        self.canvas.configure(yscrollcommand=self.v_scroll.set, xscrollcommand=self.h_scroll.set)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        self.v_scroll.grid(row=0, column=1, sticky="ns")
        self.h_scroll.grid(row=1, column=0, sticky="ew")

        self.canvas.bind("<Configure>", self._on_configure)
        self.canvas.bind("<Button-1>", self._on_click)
        self.canvas.bind("<B1-Motion>", self._on_drag)
        self.canvas.bind("<ButtonRelease-1>", self._on_release)
        self.canvas.bind("<Motion>", self._on_motion)
        self.canvas.bind("<Double-1>", self._on_double_click)
        self.canvas.bind("<MouseWheel>", self._on_mousewheel)
        self.canvas.bind("<Shift-MouseWheel>", self._on_shift_mousewheel)
        self.canvas.bind("<Leave>", self._hide_tooltip)

    def set_columns(self, headers: List[str], display_headers: Optional[List[str]] = None, redraw: bool = True) -> None:
        self.headers = list(headers)
        if display_headers is not None and len(display_headers) == len(self.headers):
            self.display_headers = [clean_cell(header) for header in display_headers]
        else:
            self.display_headers = list(self.headers)
        self.selected_index = None
        if redraw:
            self._auto_widths()
            self.redraw()

    def set_rows(self, rows: List[List[str]], row_ids: Optional[List[int]] = None, redraw: bool = True) -> None:
        self.rows = [[clean_cell(value) for value in row] for row in rows]
        self.row_ids = row_ids if row_ids is not None else list(range(len(rows)))
        if self.selected_index is not None and self.selected_index >= len(self.rows):
            self.selected_index = None
        self._auto_widths()
        if redraw:
            self.redraw()

    def set_row_states(self, confirmed_ids=None, auto_matched_ids=None, warning_ids=None) -> None:
        self.confirmed_ids = set(confirmed_ids or [])
        self.auto_matched_ids = set(auto_matched_ids or [])
        self.warning_ids = set(warning_ids or [])
        self.redraw()

    def set_confirmed_ids(self, confirmed_ids) -> None:
        self.confirmed_ids = set(confirmed_ids or [])
        self.redraw()

    def set_auto_matched_ids(self, auto_matched_ids) -> None:
        self.auto_matched_ids = set(auto_matched_ids or [])
        self.redraw()

    def set_warning_ids(self, warning_ids) -> None:
        self.warning_ids = set(warning_ids or [])
        self.redraw()

    def clear(self) -> None:
        self.rows = []
        self.row_ids = []
        self.selected_index = None
        self.redraw()

    def selection(self) -> List[str]:
        if self.selected_index is None:
            return []
        if self.selected_index < 0 or self.selected_index >= len(self.row_ids):
            return []
        return [str(self.row_ids[self.selected_index])]

    def selection_set(self, row_id: str) -> None:
        try:
            target = int(row_id)
        except ValueError:
            return
        if target in self.row_ids:
            self.selected_index = self.row_ids.index(target)
            self.redraw()

    def see(self, row_id: str) -> None:
        try:
            target = int(row_id)
        except ValueError:
            return
        if target not in self.row_ids:
            return
        index = self.row_ids.index(target)
        total_height = self.header_height + len(self.rows) * self.row_height
        visible_height = max(1, self.canvas.winfo_height())
        y = self.header_height + index * self.row_height
        current_top = self.canvas.canvasy(0)
        current_bottom = current_top + visible_height
        if y < current_top:
            self.canvas.yview_moveto(max(0, y / max(1, total_height)))
            self.redraw()
        elif y + self.row_height > current_bottom:
            self.canvas.yview_moveto(max(0, (y + self.row_height - visible_height) / max(1, total_height)))
            self.redraw()

    def center_on(self, row_id: str) -> None:
        try:
            target = int(row_id)
        except ValueError:
            return
        if target not in self.row_ids:
            return
        index = self.row_ids.index(target)
        total_height = self.header_height + len(self.rows) * self.row_height
        visible_height = max(1, self.canvas.winfo_height())
        target_y = self.header_height + index * self.row_height - max(0, (visible_height - self.row_height) // 2)
        self.canvas.yview_moveto(max(0, min(1, target_y / max(1, total_height))))
        self.redraw()

    def row_view_offset(self, row_id: str) -> int:
        try:
            target = int(row_id)
        except ValueError:
            return 0
        if target not in self.row_ids:
            return 0
        index = self.row_ids.index(target)
        row_y = self.header_height + index * self.row_height
        return int(row_y - self.canvas.canvasy(0))

    def move_row_to_view_offset(self, row_id: str, offset: int) -> None:
        self.align_row_to_offset(row_id, offset)

    def align_row_to_offset(self, row_id: str, offset: int) -> None:
        try:
            target = int(row_id)
        except ValueError:
            return
        if target not in self.row_ids:
            return
        index = self.row_ids.index(target)
        total_height = self.header_height + len(self.rows) * self.row_height
        row_y = self.header_height + index * self.row_height
        desired_top = row_y - offset
        max_top = max(0, total_height - max(1, self.canvas.winfo_height()))
        desired_top = max(0, min(max_top, desired_top))
        self.canvas.yview_moveto(desired_top / max(1, total_height))
        self.redraw()

    def scroll_row_to_offset_animated(self, row_id: str, offset: int, steps: int = 8) -> None:
        try:
            target = int(row_id)
        except ValueError:
            return
        if target not in self.row_ids:
            return
        index = self.row_ids.index(target)
        total_height = self.header_height + len(self.rows) * self.row_height
        row_y = self.header_height + index * self.row_height
        desired_top = row_y - offset
        max_top = max(0, total_height - max(1, self.canvas.winfo_height()))
        desired_top = max(0, min(max_top, desired_top))
        current_top = float(self.canvas.canvasy(0))
        if abs(desired_top - current_top) < 1:
            self.align_row_to_offset(row_id, offset)
            return
        if self.scroll_after_id is not None:
            self.after_cancel(self.scroll_after_id)
            self.scroll_after_id = None
        self._animate_scroll(current_top, desired_top, max(1, total_height), max(1, steps), 0)

    def _animate_scroll(self, start_top: float, end_top: float, total_height: int, steps: int, step: int) -> None:
        if step >= steps:
            self.canvas.yview_moveto(end_top / total_height)
            self.redraw()
            self.scroll_after_id = None
            return
        progress = (step + 1) / steps
        eased = 1 - (1 - progress) * (1 - progress)
        top = start_top + (end_top - start_top) * eased
        self.canvas.yview_moveto(top / total_height)
        self.redraw()
        self.scroll_after_id = self.after(14, lambda: self._animate_scroll(start_top, end_top, total_height, steps, step + 1))

    def align_row_to_standard_position(self, row_id: str, body_rows_from_top: int = 4) -> None:
        offset = self.header_height + body_rows_from_top * self.row_height
        self.align_row_to_offset(row_id, offset)

    def identify_cell(self, x: int, y: int) -> Tuple[Optional[int], Optional[str], Optional[Tuple[int, int, int, int]]]:
        canvas_x = int(self.canvas.canvasx(x))
        canvas_y = int(self.canvas.canvasy(y))
        # 表头固定在可视区域顶部；点击表头时不能按滚动后的画布坐标识别成数据行。
        if y < self.header_height:
            return None, None, None
        row_index = (canvas_y - self.header_height) // self.row_height
        if row_index < 0 or row_index >= len(self.rows):
            return None, None, None
        left = 0
        for col_index, width in enumerate(self.column_widths):
            right = left + width
            if left <= canvas_x < right:
                header = self.headers[col_index] if col_index < len(self.headers) else ""
                return int(row_index), header, (left, self.header_height + int(row_index) * self.row_height, width, self.row_height)
            left = right
        return None, None, None

    def get_value(self, row_index: int, header: str) -> str:
        if header not in self.headers or row_index < 0 or row_index >= len(self.rows):
            return ""
        col_index = self.headers.index(header)
        if col_index >= len(self.rows[row_index]):
            return ""
        return self.rows[row_index][col_index]

    def place_editor(self, row_index: int, header: str, editor: tk.Widget) -> None:
        if header not in self.headers:
            return
        col_index = self.headers.index(header)
        x = sum(self.column_widths[:col_index]) - int(self.canvas.canvasx(0))
        y = self.header_height + row_index * self.row_height - int(self.canvas.canvasy(0))
        width = self.column_widths[col_index]
        editor.place(in_=self.canvas, x=x + 1, y=y + 1, width=max(30, width - 2), height=self.row_height - 2)

    def _on_y_scroll(self, *args) -> None:
        self.canvas.yview(*args)
        self.redraw()

    def _on_x_scroll(self, *args) -> None:
        self.canvas.xview(*args)
        self.redraw()

    def _on_configure(self, _event=None) -> None:
        if self._resize_after:
            self.after_cancel(self._resize_after)
        self._resize_after = self.after(80, self._resize_redraw)

    def _resize_redraw(self) -> None:
        self._resize_after = None
        self._auto_widths()
        self.redraw()

    def _on_click(self, event) -> None:
        resize_col = self._header_resize_hit(event.x, event.y)
        if resize_col is not None:
            self.resizing_col = resize_col
            self.resize_start_x = int(self.canvas.canvasx(event.x))
            self.resize_start_width = self.column_widths[resize_col]
            return
        row_index, header, _bbox = self.identify_cell(event.x, event.y)
        if row_index is None:
            return
        row_id = self.row_ids[row_index]
        if header == "确认" and self.on_checkbox:
            self.on_checkbox(row_id)
            return
        self.selected_index = row_index
        self.redraw()
        if self.on_select:
            self.on_select(row_id)

    def _on_drag(self, event) -> None:
        if self.resizing_col is None:
            return
        current_x = int(self.canvas.canvasx(event.x))
        delta = current_x - self.resize_start_x
        new_width = max(24, self.resize_start_width + delta)
        self.column_widths[self.resizing_col] = new_width
        if self.resizing_col < len(self.headers):
            self.manual_widths[self.headers[self.resizing_col]] = new_width
        self.redraw()

    def _on_release(self, _event) -> None:
        self.resizing_col = None

    def _on_motion(self, event) -> None:
        if self._header_resize_hit(event.x, event.y) is not None:
            self.canvas.configure(cursor="sb_h_double_arrow")
        else:
            self.canvas.configure(cursor="")
        self._schedule_tooltip(event)

    def _schedule_tooltip(self, event) -> None:
        row_index, header, _bbox = self.identify_cell(event.x, event.y)
        cell = (row_index, header) if row_index is not None and header else None
        if cell == self.tooltip_cell:
            return
        self.tooltip_cell = cell
        self._hide_tooltip()
        if cell is None or header == "确认":
            return
        value = self.get_value(row_index, header)
        if not value:
            return
        screen_x = self.canvas.winfo_rootx() + event.x + 16
        screen_y = self.canvas.winfo_rooty() + event.y + 14
        self.tooltip_after = self.after(1000, lambda: self._show_tooltip(value, screen_x, screen_y))

    def _show_tooltip(self, text: str, screen_x: int, screen_y: int) -> None:
        self._hide_tooltip(cancel_after=False)
        tooltip = tk.Toplevel(self)
        tooltip.wm_overrideredirect(True)
        tooltip.wm_geometry(f"+{screen_x}+{screen_y}")
        label = tk.Label(
            tooltip,
            text=text,
            justify="left",
            background="#FFF8C6",
            foreground="#111111",
            relief="solid",
            borderwidth=1,
            padx=8,
            pady=5,
            wraplength=680,
            font=("Microsoft YaHei UI", 10),
        )
        label.pack()
        self.tooltip_window = tooltip

    def _hide_tooltip(self, _event=None, cancel_after: bool = True) -> None:
        if cancel_after and self.tooltip_after:
            self.after_cancel(self.tooltip_after)
            self.tooltip_after = None
        if self.tooltip_window is not None:
            self.tooltip_window.destroy()
            self.tooltip_window = None

    def _header_resize_hit(self, x: int, y: int) -> Optional[int]:
        # 使用窗口坐标判断固定表头，纵向滚动后仍可拖动每个分隔线调整列宽。
        if y < 0 or y > self.header_height:
            return None
        canvas_x = int(self.canvas.canvasx(x))
        left = 0
        for index, width in enumerate(self.column_widths):
            right = left + width
            if abs(canvas_x - right) <= 7:
                return index
            left = right
        return None

    def _on_double_click(self, event) -> None:
        row_index, header, bbox = self.identify_cell(event.x, event.y)
        if row_index is None or header is None:
            return
        self.selected_index = row_index
        self.redraw()
        if self.on_double_click:
            self.on_double_click(row_index, header, bbox)

    def _on_mousewheel(self, event) -> None:
        self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
        self.redraw()

    def _on_shift_mousewheel(self, event) -> None:
        self.canvas.xview_scroll(int(-1 * (event.delta / 120)), "units")
        self.redraw()

    def _auto_widths(self) -> None:
        if not self.headers:
            self.column_widths = []
            return
        available = max(400, self.canvas.winfo_width() - 4)
        sample_rows = self.rows[:80]
        natural: List[int] = []
        weights: List[float] = []
        for col_index, header in enumerate(self.headers):
            display_header = self.display_headers[col_index] if col_index < len(self.display_headers) else header
            values = [display_header] + [row[col_index] if col_index < len(row) else "" for row in sample_rows]
            if header in self.manual_widths:
                preferred = self.manual_widths[header]
            else:
                preferred = preferred_column_width(header, values)
            natural.append(preferred)
            weights.append(column_weight(header))

        total = sum(natural)
        if total < available and natural:
            extra = available - total if self.fill_available_width else min(available - total, int(total * 0.35))
            weight_total = sum(weights)
            natural = [width + int(extra * weights[index] / weight_total) for index, width in enumerate(natural)]
        self.column_widths = natural

    def redraw(self) -> None:
        self.canvas.delete("all")
        total_width = max(sum(self.column_widths), self.canvas.winfo_width())
        total_height = self.header_height + len(self.rows) * self.row_height
        self.canvas.configure(scrollregion=(0, 0, total_width, total_height))
        if not self.headers:
            return

        for row_index, row in enumerate(self.rows):
            y = self.header_height + row_index * self.row_height
            row_id = self.row_ids[row_index] if row_index < len(self.row_ids) else row_index
            fill = self.auto_matched_bg if row_id in self.auto_matched_ids else "white"
            if row_id in self.warning_ids:
                fill = self.warning_bg
            if row_id in self.confirmed_ids:
                fill = self.confirmed_bg
            if row_index == self.selected_index:
                fill = self.selected_bg
            x = 0
            for col_index, header in enumerate(self.headers):
                width = self.column_widths[col_index]
                self.canvas.create_rectangle(x, y, x + width, y + self.row_height, fill=fill, outline=self.grid_color)
                value = row[col_index] if col_index < len(row) else ""
                if header == "确认":
                    self._draw_checkbox(x, y, width, self.row_height, value in ["1", "是", "TRUE", "true", "✓"])
                else:
                    max_lines = self.visible_text_lines(self.font)
                    self._draw_text(value, x + 3, y + 1, width - 6, self.row_height - 2, self.font, max_lines=max_lines, force_wrap=is_long_text_column(header) and max_lines > 1)
                x += width

        # 最后绘制表头，并将它放在当前可视区域顶端，使纵向滚动只移动数据行。
        header_top = int(self.canvas.canvasy(0))
        x = 0
        for col_index, header in enumerate(self.headers):
            width = self.column_widths[col_index]
            display_header = self.display_headers[col_index] if col_index < len(self.display_headers) else header
            self.canvas.create_rectangle(
                x,
                header_top,
                x + width,
                header_top + self.header_height,
                fill=self.header_bg,
                outline=self.grid_color,
            )
            self._draw_text(
                display_header,
                x + 4,
                header_top + 3,
                width - 8,
                self.header_height - 6,
                self.header_font,
                center=True,
                max_lines=1,
            )
            x += width

    def _draw_checkbox(self, x: int, y: int, width: int, height: int, checked: bool) -> None:
        size = min(13, max(10, height - 7))
        left = x + (width - size) // 2
        top = y + (height - size) // 2
        self.canvas.create_rectangle(left, top, left + size, top + size, fill="white", outline="#666666")
        if checked:
            self.canvas.create_line(left + 3, top + size // 2, left + size // 2 - 1, top + size - 4, left + size - 3, top + 3, fill="#16803A", width=2)

    def _draw_text(self, text: str, x: int, y: int, width: int, height: int, font: tkfont.Font, center: bool = False, max_lines: int = 2, force_wrap: bool = False) -> None:
        if width <= 4:
            return
        lines = wrap_text_to_lines(clean_cell(text), font, max(8, width), max_lines, force_wrap=force_wrap)
        line_height = font.metrics("linespace")
        total_height = min(len(lines), max_lines) * line_height
        current_y = y + max(0, (height - total_height) // 2)
        anchor = "n" if center else "nw"
        text_x = x + width // 2 if center else x
        for line in lines[:max_lines]:
            self.canvas.create_text(text_x, current_y, text=line, anchor=anchor, font=font, fill=self.text_color)
            current_y += line_height

    def visible_text_lines(self, font: tkfont.Font) -> int:
        line_height = max(1, font.metrics("linespace"))
        available = max(1, self.row_height - 4)
        return max(1, min(self.max_lines, available // line_height))


def text_display_width(value: str) -> int:
    total = 0
    for char in clean_cell(value):
        total += 2 if ord(char) > 127 else 1
    return total


def percentile(values: List[int], ratio: float) -> int:
    if not values:
        return 0
    sorted_values = sorted(values)
    index = int((len(sorted_values) - 1) * ratio)
    return sorted_values[index]


def preferred_column_width(header: str, values: List[str]) -> int:
    key = normalize_key(header)
    if key == normalize_key("确认"):
        return 28
    if is_small_numeric_column(header):
        return 34
    if is_code_column(header):
        sample = percentile([text_display_width(value) for value in values], 0.65)
        return min(max(sample * 7 + 14, 76), 110)
    if is_long_text_column(header):
        sample = percentile([text_display_width(value) for value in values], 0.55)
        return min(max(sample * 7 + 14, 92), 180)
    sample = percentile([text_display_width(value) for value in values], 0.65)
    return min(max(sample * 7 + 14, 54), 130)


def column_weight(header: str) -> float:
    if is_small_numeric_column(header) or normalize_key(header) == normalize_key("确认"):
        return 0.05
    if is_code_column(header):
        return 0.55
    if is_long_text_column(header):
        return 0.8
    return 0.45


def is_small_numeric_column(header: str) -> bool:
    text = normalize_key(header)
    keywords = ["层级", "工序", "总用量", "用量", "数量", "点数", "备注", "quantity", "qty", "count", "匹配分数", "单位", "unit"]
    return any(normalize_key(keyword) == text or normalize_key(keyword) in text for keyword in keywords)


def is_code_column(header: str) -> bool:
    text = normalize_key(header)
    keywords = ["物料编码", "编码", "code", "partnumber", "p/n"]
    return any(normalize_key(keyword) in text for keyword in keywords)


def is_long_text_column(header: str) -> bool:
    text = normalize_key(header)
    keywords = ["规格", "型号", "描述", "备注", "位号", "物料名称", "comment", "description", "spec", "designator", "footprint", "libref", "name"]
    return any(normalize_key(keyword) in text for keyword in keywords)


def wrap_text_to_lines(text: str, font: tkfont.Font, width: int, max_lines: int, force_wrap: bool = False) -> List[str]:
    text = clean_cell(text)
    if not text:
        return [""]
    if force_wrap and max_lines > 1 and should_force_wrap(text, font, width):
        forced = force_split_text(text, font, width, max_lines)
        if forced:
            return forced
    tokens = split_wrap_tokens(text)
    lines: List[str] = []
    current = ""
    for token in tokens:
        candidate = current + token
        if current and font.measure(candidate) > width:
            lines.append(current)
            current = token.lstrip()
            if len(lines) == max_lines:
                break
        else:
            current = candidate
    if len(lines) < max_lines and current:
        lines.append(current)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
    if lines and font.measure(lines[-1]) > width:
        lines[-1] = ellipsize(lines[-1], font, width)
    if len(lines) == max_lines and "".join(lines) != text:
        lines[-1] = ellipsize(lines[-1], font, width)
    return lines or [""]


def should_force_wrap(text: str, font: tkfont.Font, width: int) -> bool:
    if font.measure(text) > width:
        return True
    if len(text) >= 14 and re.search(r"[/,，;；、\s]", text):
        return True
    return False


def force_split_text(text: str, font: tkfont.Font, width: int, max_lines: int) -> List[str]:
    if max_lines < 2:
        return [ellipsize(text, font, width)]
    split_points = [match.end() for match in re.finditer(r"[/,，;；、\s]", text)]
    if not split_points:
        split_points = list(range(8, len(text), 8))
    target = max(1, len(text) // 2)
    split_points.sort(key=lambda point: abs(point - target))
    for point in split_points:
        first = text[:point].strip()
        second = text[point:].strip()
        if not first or not second:
            continue
        if font.measure(first) <= width and font.measure(second) <= width:
            return [first, second]
    lines = wrap_text_to_lines(text, font, width, max_lines, force_wrap=False)
    return lines


def split_wrap_tokens(text: str) -> List[str]:
    parts = re.findall(r"[A-Za-z0-9.%+\-]+|[/,，;；、\s]+|[\u4e00-\u9fff]|[^\w\u4e00-\u9fff]", text)
    return parts if parts else [text]


def ellipsize(text: str, font: tkfont.Font, width: int) -> str:
    if font.measure(text) <= width:
        return text
    suffix = "..."
    result = text
    while result and font.measure(result + suffix) > width:
        result = result[:-1]
    return (result + suffix) if result else suffix




def draw_rounded_rect(canvas: tk.Canvas, x1: int, y1: int, x2: int, y2: int, radius: int, **kwargs) -> int:
    width = max(1, int(round(x2 - x1)))
    height = max(1, int(round(y2 - y1)))
    radius = max(0, min(int(radius), width // 2, height // 2))
    fill = kwargs.pop("fill", "")
    outline = kwargs.pop("outline", "")
    line_width = max(1, int(kwargs.pop("width", 1)))
    tags = kwargs.pop("tags", None)

    # Tk Canvas 的圆弧在 Windows 上没有抗锯齿。这里按 2~4 倍分辨率实时绘制，
    # 再用 Lanczos 缩小成透明图片，得到平滑边缘；不是外部图片或抠图。
    area = width * height
    scale = 4 if area <= 24_000 else 3 if area <= 140_000 else 2
    high_width = width * scale
    high_height = height * scale
    high_image = Image.new("RGBA", (high_width, high_height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(high_image)
    draw.rounded_rectangle(
        (0, 0, high_width - 1, high_height - 1),
        radius=radius * scale,
        fill=fill or None,
        outline=outline or None,
        width=line_width * scale,
    )
    smooth_image = high_image.resize((width, height), Image.Resampling.LANCZOS)
    photo = ImageTk.PhotoImage(smooth_image, master=canvas)

    image_refs = getattr(canvas, "_aa_image_refs", {})
    live_items = set(canvas.find_all())
    for item_id in list(image_refs):
        if item_id not in live_items:
            del image_refs[item_id]
    create_kwargs = {"tags": tags} if tags is not None else {}
    image_item = canvas.create_image(int(round(x1)), int(round(y1)), anchor="nw", image=photo, **create_kwargs)
    image_refs[image_item] = photo
    canvas._aa_image_refs = image_refs
    return image_item


class RoundedCard(tk.Frame):
    def __init__(self, parent, bg: str = "#FFFFFF", outer_bg: str = "#F5F5F7", radius: int = 24, border: str = "#E5E5EA", padding: int = 12):
        super().__init__(parent, bg=outer_bg)
        self.card_bg = bg
        self.outer_bg = outer_bg
        self.radius = radius
        self.border = border
        self.padding = padding
        self.canvas = tk.Canvas(self, bg=outer_bg, highlightthickness=0, bd=0)
        self.canvas.pack(fill="both", expand=True)
        self.body = tk.Frame(self.canvas, bg=bg)
        self.window_id = self.canvas.create_window(padding, padding, anchor="nw", window=self.body)
        self.canvas.bind("<Configure>", self._redraw)

    def _redraw(self, event=None) -> None:
        width = max(1, self.canvas.winfo_width())
        height = max(1, self.canvas.winfo_height())
        self.canvas.delete("card")
        draw_rounded_rect(
            self.canvas,
            1,
            1,
            width - 1,
            height - 1,
            self.radius,
            fill=self.card_bg,
            outline=self.border,
            width=1,
            tags="card",
        )
        self.canvas.tag_lower("card")
        inner_width = max(1, width - self.padding * 2)
        inner_height = max(1, height - self.padding * 2)
        self.canvas.coords(self.window_id, self.padding, self.padding)
        self.canvas.itemconfigure(self.window_id, width=inner_width, height=inner_height)


class AppleButton(tk.Canvas):
    PALETTES = {
        "primary": ("#0A84FF", "#0077ED", "#006EDB", "#FFFFFF", "#0A84FF"),
        "secondary": ("#EEF5FF", "#E2EFFF", "#D6E8FF", "#0A63CE", "#D2E5FA"),
        "neutral": ("#FFFFFF", "#F2F2F7", "#E9E9EF", "#1C1C1E", "#D1D1D6"),
        "selected": ("#DCEBFF", "#D2E5FC", "#C7DDF8", "#0A63CE", "#8CB9EE"),
    }

    def __init__(
        self,
        parent,
        text: str,
        command: Optional[Callable[[], None]] = None,
        variant: str = "neutral",
        width: int = 120,
        height: int = 42,
        enabled: bool = True,
        canvas_bg: Optional[str] = None,
    ):
        if canvas_bg is None:
            try:
                canvas_bg = parent.cget("background")
            except tk.TclError:
                canvas_bg = "#F5F5F3"
        super().__init__(parent, width=width, height=height, bg=canvas_bg or "#F5F5F3", highlightthickness=0, bd=0)
        self.button_text = text
        self.command = command
        self.variant = variant
        self.enabled = bool(enabled)
        self.button_width = width
        self.button_height = height
        self.hovered = False
        self.pressed = False
        self.configure(cursor="hand2" if self.enabled else "arrow")
        self.bind("<Configure>", self._redraw)
        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        self.bind("<ButtonPress-1>", self._on_press)
        self.bind("<ButtonRelease-1>", self._on_release)
        self._redraw()

    def _on_enter(self, _event=None) -> None:
        if self.enabled:
            self.hovered = True
            self._redraw()

    def _on_leave(self, _event=None) -> None:
        self.hovered = False
        self.pressed = False
        self._redraw()

    def _on_press(self, _event=None) -> None:
        if self.enabled:
            self.pressed = True
            self._redraw()

    def _on_release(self, event=None) -> None:
        if not self.enabled:
            return
        was_pressed = self.pressed
        self.pressed = False
        inside = event is None or (0 <= event.x < self.winfo_width() and 0 <= event.y < self.winfo_height())
        self._redraw()
        if was_pressed and inside and self.command is not None:
            self.command()

    def set_enabled(self, enabled: bool) -> None:
        self.enabled = bool(enabled)
        self.hovered = False
        self.pressed = False
        self.configure(cursor="hand2" if self.enabled else "arrow")
        self._redraw()

    def set_variant(self, variant: str) -> None:
        if variant in self.PALETTES:
            self.variant = variant
            self._redraw()

    def _redraw(self, _event=None) -> None:
        self.delete("all")
        width = max(self.button_width, self.winfo_width())
        height = max(self.button_height, self.winfo_height())
        if self.enabled:
            normal, hover, pressed, foreground, border = self.PALETTES.get(self.variant, self.PALETTES["neutral"])
            fill = pressed if self.pressed else hover if self.hovered else normal
        else:
            fill, foreground, border = "#E5E5EA", "#A1A1A6", "#D8D8DC"
        draw_rounded_rect(
            self,
            1,
            1,
            width - 1,
            height - 1,
            min(16, height // 2),
            fill=fill,
            outline=border,
            width=1,
        )
        self.create_text(
            width // 2,
            height // 2,
            text=self.button_text,
            fill=foreground,
            font=("Microsoft YaHei UI", 10, "bold" if self.variant in ("primary", "selected") else "normal"),
            anchor="center",
        )


class AppleNavItem(tk.Canvas):
    def __init__(self, parent, text: str, icon_text: str, icon_color: str, command: Callable[[], None]):
        super().__init__(parent, height=56, bg="#FFFFFF", highlightthickness=0, bd=0, cursor="hand2")
        self.item_text = text
        self.icon_text = icon_text
        self.icon_color = icon_color
        self.command = command
        self.selected = False
        self.hovered = False
        self.pressed = False
        self.bind("<Configure>", self._redraw)
        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        self.bind("<ButtonPress-1>", self._on_press)
        self.bind("<ButtonRelease-1>", self._on_release)

    def set_selected(self, selected: bool) -> None:
        self.selected = bool(selected)
        self.pressed = False
        self._redraw()

    def _on_enter(self, _event=None) -> None:
        self.hovered = True
        self._redraw()

    def _on_leave(self, _event=None) -> None:
        self.hovered = False
        self.pressed = False
        self._redraw()

    def _on_press(self, _event=None) -> None:
        self.pressed = True
        self._redraw()

    def _on_release(self, event=None) -> None:
        was_pressed = self.pressed
        self.pressed = False
        inside = event is None or (0 <= event.x < self.winfo_width() and 0 <= event.y < self.winfo_height())
        self._redraw()
        if was_pressed and inside:
            self.command()

    def _redraw(self, _event=None) -> None:
        self.delete("all")
        width = max(240, self.winfo_width())
        height = max(56, self.winfo_height())
        if self.selected:
            fill = "#789BB9" if not self.pressed else "#6F90AC"
            foreground = "#FFFFFF"
            chevron_color = "#FFFFFF"
        else:
            fill = "#E7E9ED" if self.pressed else "#EEF0F4" if self.hovered else "#F8F8FA"
            foreground = "#1C1C1E"
            chevron_color = "#A9A9AE"
        draw_rounded_rect(
            self,
            4,
            3,
            width - 4,
            height - 3,
            18,
            fill=fill,
            outline="" if self.selected else "#ECECF0",
            width=1,
        )
        draw_rounded_rect(self, 15, 12, 45, 42, 8, fill=self.icon_color, outline="")
        self.create_text(30, 27, text=self.icon_text, fill="#FFFFFF", font=("Microsoft YaHei UI", 10, "bold"), anchor="center")
        self.create_text(58, height // 2, text=self.item_text, fill=foreground, font=("Microsoft YaHei UI", 11, "bold" if self.selected else "normal"), anchor="w")
        self.create_text(width - 20, height // 2 - 1, text="›", fill=chevron_color, font=("Microsoft YaHei UI", 20), anchor="center")


class IosSwitch(tk.Canvas):
    def __init__(
        self,
        parent,
        variable: Optional[tk.BooleanVar] = None,
        command: Optional[Callable[[], None]] = None,
        width: int = 54,
        height: int = 32,
    ):
        self.switch_width = max(48, width)
        self.switch_height = max(28, height)
        super().__init__(
            parent,
            width=self.switch_width,
            height=self.switch_height,
            bg=parent.cget("bg"),
            highlightthickness=0,
            bd=0,
            cursor="hand2",
        )
        self.variable = variable or tk.BooleanVar(value=False)
        self.command = command
        self._trace_name = self.variable.trace_add("write", self._on_variable_changed)
        self.bind("<Button-1>", self._toggle)
        self._redraw()

    def _toggle(self, _event=None) -> None:
        self.variable.set(not bool(self.variable.get()))
        if self.command is not None:
            self.command()

    def _on_variable_changed(self, *_args) -> None:
        if self.winfo_exists():
            self._redraw()

    def _redraw(self) -> None:
        self.delete("all")
        enabled = bool(self.variable.get())
        margin = 2
        track_height = self.switch_height - margin * 2
        knob_size = track_height - 4
        draw_rounded_rect(
            self,
            1,
            margin,
            self.switch_width - 1,
            self.switch_height - margin,
            track_height // 2,
            fill="#34C759" if enabled else "#E5E5EA",
            outline="",
        )
        knob_left = self.switch_width - margin - knob_size - 2 if enabled else margin + 2
        knob_top = margin + 2
        draw_rounded_rect(
            self,
            knob_left,
            knob_top,
            knob_left + knob_size,
            knob_top + knob_size,
            knob_size // 2,
            fill="#FFFFFF",
            outline="#D1D1D6",
            width=1,
        )

    def destroy(self) -> None:
        try:
            self.variable.trace_remove("write", self._trace_name)
        except (tk.TclError, AttributeError):
            pass
        super().destroy()


class MobileVerticalScrollbar(tk.Canvas):
    """A slim, trackless vertical scrollbar styled like a mobile scroll indicator."""

    def __init__(self, parent, command: Callable[..., None], bg: str = "#F5F5F3", width: int = 11):
        super().__init__(
            parent,
            width=width,
            bg=bg,
            highlightthickness=0,
            bd=0,
            cursor="hand2",
            takefocus=0,
        )
        self.command = command
        self.first = 0.0
        self.last = 1.0
        self.hovered = False
        self.dragging = False
        self.drag_start_y = 0
        self.drag_start_first = 0.0
        self.thumb_top = 0
        self.thumb_bottom = 0
        self.bind("<Configure>", self._redraw)
        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        self.bind("<ButtonPress-1>", self._on_press)
        self.bind("<B1-Motion>", self._on_drag)
        self.bind("<ButtonRelease-1>", self._on_release)

    def set(self, first: Any, last: Any) -> None:
        try:
            self.first = max(0.0, min(1.0, float(first)))
            self.last = max(self.first, min(1.0, float(last)))
        except (TypeError, ValueError):
            self.first, self.last = 0.0, 1.0
        self._redraw()

    def _thumb_geometry(self) -> Tuple[int, int]:
        height = max(1, self.winfo_height())
        track_top = 7
        track_height = max(1, height - track_top * 2)
        visible_fraction = max(0.0, min(1.0, self.last - self.first))
        thumb_height = max(36, int(round(track_height * visible_fraction)))
        thumb_height = min(track_height, thumb_height)
        travel = max(0, track_height - thumb_height)
        scrollable = max(0.000001, 1.0 - visible_fraction)
        thumb_top = track_top + int(round(travel * self.first / scrollable)) if travel else track_top
        return thumb_top, thumb_top + thumb_height

    def _redraw(self, _event=None) -> None:
        if not self.winfo_exists():
            return
        self.delete("all")
        if self.last - self.first >= 0.999:
            self.thumb_top = self.thumb_bottom = 0
            return
        self.thumb_top, self.thumb_bottom = self._thumb_geometry()
        widget_width = max(7, self.winfo_width())
        thumb_width = 6 if self.hovered or self.dragging else 4
        left = (widget_width - thumb_width) // 2
        draw_rounded_rect(
            self,
            left,
            self.thumb_top,
            left + thumb_width,
            self.thumb_bottom,
            thumb_width // 2,
            fill="#78787E" if self.dragging else "#99999F",
            outline="",
            tags="thumb",
        )

    def _on_enter(self, _event=None) -> None:
        self.hovered = True
        self._redraw()

    def _on_leave(self, _event=None) -> None:
        if not self.dragging:
            self.hovered = False
            self._redraw()

    def _on_press(self, event) -> None:
        if self.last - self.first >= 0.999:
            return
        if not (self.thumb_top <= event.y <= self.thumb_bottom):
            direction = -1 if event.y < self.thumb_top else 1
            self.command("scroll", direction, "pages")
            return
        self.dragging = True
        self.drag_start_y = event.y
        self.drag_start_first = self.first
        self._redraw()

    def _on_drag(self, event) -> None:
        if not self.dragging:
            return
        height = max(1, self.winfo_height())
        track_height = max(1, height - 14)
        thumb_height = max(1, self.thumb_bottom - self.thumb_top)
        travel = max(1, track_height - thumb_height)
        visible_fraction = max(0.0, min(1.0, self.last - self.first))
        new_first = self.drag_start_first + (event.y - self.drag_start_y) / travel * (1.0 - visible_fraction)
        self.command("moveto", max(0.0, min(1.0 - visible_fraction, new_first)))

    def _on_release(self, _event=None) -> None:
        self.dragging = False
        self._redraw()


class ColumnRuleEditor(tk.Frame):
    MIN_CARD_WIDTH = 68
    MAX_CARD_WIDTH = 112
    CARD_HEIGHT = 150
    CARD_GAP = 3

    def __init__(self, parent, order: List[str], custom_titles: Dict[str, str], enabled: bool = False):
        super().__init__(parent, bg="#FFFFFF")
        self.order = normalize_result_column_order(order)
        self.title_vars = {
            key: tk.StringVar(value=custom_titles.get(key, ""))
            for key in CUSTOMIZABLE_RESULT_COLUMN_KEYS
        }
        self.card_width = self.MAX_CARD_WIDTH
        self.enabled = enabled
        self.drag_key: Optional[str] = None
        self.drag_drop_index: Optional[int] = None
        self.drag_finishing = False
        self.drag_motion_bind_id: Optional[str] = None
        self.drag_release_bind_id: Optional[str] = None
        self.column_frames: Dict[str, tk.Frame] = {}

        self.columnconfigure(1, weight=1)
        labels = tk.Frame(self, bg="#FFFFFF", width=66, height=self.CARD_HEIGHT)
        labels.grid(row=0, column=0, sticky="ns", padx=(0, 6))
        labels.grid_propagate(False)
        for row, (text, height) in enumerate((("默认列名", 64), ("替代名称", 48), ("拖动排序", 38))):
            labels.rowconfigure(row, minsize=height)
            tk.Label(
                labels,
                text=text,
                bg="#FFFFFF",
                fg="#1C1C1E" if row < 2 else "#8E8E93",
                font=("Microsoft YaHei UI", 10, "bold" if row < 2 else "normal"),
                anchor="w",
            ).grid(row=row, column=0, sticky="nsew")

        self.canvas = tk.Canvas(self, bg="#FFFFFF", height=self.CARD_HEIGHT, highlightthickness=0, bd=0)
        self.canvas.grid(row=0, column=1, sticky="ew")
        self.scrollbar = ttk.Scrollbar(self, orient="horizontal", command=self.canvas.xview)
        self.scrollbar.grid(row=1, column=1, sticky="ew", pady=(6, 0))
        self.canvas.configure(xscrollcommand=self.scrollbar.set)
        self.inner = tk.Frame(self.canvas, bg="#FFFFFF")
        self.inner_window = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", self._update_scrollregion)
        self.canvas.bind("<Configure>", self._resize_cards_to_fit)
        self.canvas.bind("<Shift-MouseWheel>", self._on_shift_mousewheel)

        self.drop_indicator_width = 8
        self.drop_indicator = tk.Frame(
            self,
            bg="#0A84FF",
            width=self.drop_indicator_width,
            height=self.CARD_HEIGHT - 24,
        )
        self.drop_indicator.place_forget()
        self.drop_badge_width = 80
        self.drop_badge_height = 28
        self.drop_badge_font = tkfont.Font(family="Microsoft YaHei UI", size=9, weight="bold")
        self.drop_badge = tk.Canvas(
            self,
            width=self.drop_badge_width,
            height=self.drop_badge_height,
            bg="#FFFFFF",
            highlightthickness=0,
            bd=0,
        )
        self.drop_badge_text_id: Optional[int] = None
        self._set_drop_badge_text("移动列")
        self.drop_badge.place_forget()
        self._render_columns()

    def _update_scrollregion(self, _event=None) -> None:
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        content_width = self.inner.winfo_reqwidth()
        if content_width <= max(1, self.canvas.winfo_width()):
            self.scrollbar.grid_remove()
            self.canvas.xview_moveto(0)
        else:
            self.scrollbar.grid()

    def _on_shift_mousewheel(self, event) -> None:
        self.canvas.xview_scroll(-int(event.delta / 120), "units")

    def _resize_cards_to_fit(self, event) -> None:
        if not self.order:
            return
        gaps = self.CARD_GAP * max(0, len(self.order) - 1)
        fitted_width = (max(1, event.width) - gaps) // len(self.order)
        target_width = max(self.MIN_CARD_WIDTH, min(self.MAX_CARD_WIDTH, fitted_width))
        if target_width == self.card_width:
            return
        self.card_width = target_width
        for card in self.column_frames.values():
            card.configure(width=self.card_width)
            default_label = getattr(card, "default_label", None)
            if default_label is not None:
                default_label.configure(wraplength=0)
        self.inner.update_idletasks()
        self._update_scrollregion()

    def _render_columns(self) -> None:
        view_start = self.canvas.xview()[0] if self.canvas.winfo_exists() else 0.0
        for child in self.inner.winfo_children():
            child.destroy()
        self.column_frames = {}
        for index, key in enumerate(self.order):
            card = tk.Frame(
                self.inner,
                bg="#FAFAFC" if self.enabled else "#F3F3F5",
                width=self.card_width,
                height=self.CARD_HEIGHT,
                highlightbackground="#E1E1E6",
                highlightthickness=1,
            )
            card.grid(row=0, column=index, padx=(0, self.CARD_GAP if index < len(self.order) - 1 else 0), sticky="ns")
            card.grid_propagate(False)
            card.columnconfigure(0, weight=1)
            card.rowconfigure(0, minsize=64)
            card.rowconfigure(1, minsize=48)
            card.rowconfigure(2, minsize=38)
            card_bg = card.cget("bg")
            default_label = tk.Label(
                card,
                text=default_result_column_title(key),
                bg=card_bg,
                fg="#1C1C1E",
                font=("Microsoft YaHei UI", 10, "bold"),
                justify="center",
                wraplength=0,
            )
            default_label.grid(row=0, column=0, sticky="nsew", padx=3)
            card.default_label = default_label
            entry = tk.Entry(
                card,
                textvariable=self.title_vars[key],
                relief="flat",
                bd=0,
                bg="#FFFFFF",
                disabledbackground="#E7E7EA",
                disabledforeground="#9A9AA0",
                fg="#1C1C1E",
                highlightthickness=1,
                highlightbackground="#D5D5DA",
                highlightcolor="#4F86C6",
                font=("Microsoft YaHei UI", 10),
                justify="center",
                state="normal" if self.enabled else "disabled",
            )
            entry.grid(row=1, column=0, sticky="ew", padx=3, pady=8)
            handle = tk.Label(
                card,
                text="⋮⋮ 拖动",
                bg=card_bg,
                fg="#4F86C6" if self.enabled else "#A8A8AD",
                font=("Microsoft YaHei UI", 9),
                cursor="fleur" if self.enabled else "arrow",
            )
            handle.grid(row=2, column=0, sticky="nsew")
            card.drag_handle = handle
            if self.enabled:
                handle.bind("<ButtonPress-1>", lambda event, column_key=key: self._start_drag(event, column_key))
            self.column_frames[key] = card
        self.inner.update_idletasks()
        self._update_scrollregion()
        self.after_idle(lambda: self.canvas.xview_moveto(view_start))

    def _start_drag(self, event, key: str) -> None:
        if not self.enabled or self.drag_key is not None or self.drag_finishing:
            return
        self.drag_key = key
        # Do not replace Tk's implicit mouse-button grab.  Listen at the root
        # bindtag while dragging so motion/release is still observed when the
        # pointer leaves the small handle, then remove only our callbacks.
        root = self.winfo_toplevel()
        self.drag_motion_bind_id = root.bind("<B1-Motion>", self._drag_motion, add="+")
        self.drag_release_bind_id = root.bind("<ButtonRelease-1>", self._finish_drag, add="+")
        self._set_drop_badge_text(self._drag_display_title(key))
        dragged_card = self.column_frames.get(key)
        if dragged_card is not None:
            lifted_bg = "#EDF6FF"
            dragged_card.configure(
                bg=lifted_bg,
                highlightbackground="#0A84FF",
                highlightthickness=2,
            )
            for child in dragged_card.winfo_children():
                if isinstance(child, tk.Label):
                    child.configure(bg=lifted_bg)
            handle = getattr(dragged_card, "drag_handle", None)
            if handle is not None:
                handle.configure(text="正在移动", fg="#0A84FF", font=("Microsoft YaHei UI", 9, "bold"))
        self.drag_drop_index = self._insertion_index(event.x_root)
        self._show_drop_indicator(self.drag_drop_index)

    def _remaining_drag_keys(self) -> List[str]:
        return [key for key in self.order if key != self.drag_key]

    def _drag_display_title(self, key: str) -> str:
        custom_title = self.title_vars[key].get().strip()
        return custom_title or default_result_column_title(key)

    def _set_drop_badge_text(self, text: str) -> None:
        self.drop_badge_width = max(72, min(260, self.drop_badge_font.measure(text) + 28))
        self.drop_badge.configure(width=self.drop_badge_width, height=self.drop_badge_height)
        self.drop_badge.delete("all")
        draw_rounded_rect(
            self.drop_badge,
            1,
            1,
            self.drop_badge_width - 1,
            self.drop_badge_height - 1,
            12,
            fill="#0A84FF",
            outline="",
        )
        self.drop_badge_text_id = self.drop_badge.create_text(
            self.drop_badge_width // 2,
            self.drop_badge_height // 2,
            text=text,
            fill="#FFFFFF",
            font=self.drop_badge_font,
        )

    def _insertion_index(self, pointer_x: int) -> int:
        centers = [
            self.column_frames[key].winfo_rootx() + self.column_frames[key].winfo_width() // 2
            for key in self._remaining_drag_keys()
            if key in self.column_frames and self.column_frames[key].winfo_ismapped()
        ]
        return sum(pointer_x > center for center in centers)

    def _show_drop_indicator(self, insertion_index: int) -> None:
        remaining = self._remaining_drag_keys()
        if not remaining:
            self.drop_indicator.place_forget()
            self.drop_badge.place_forget()
            return
        insertion_index = max(0, min(insertion_index, len(remaining)))
        if insertion_index == 0:
            boundary_root_x = self.column_frames[remaining[0]].winfo_rootx() - self.CARD_GAP // 2
        elif insertion_index == len(remaining):
            last_card = self.column_frames[remaining[-1]]
            boundary_root_x = last_card.winfo_rootx() + last_card.winfo_width() + self.CARD_GAP // 2
        else:
            left_card = self.column_frames[remaining[insertion_index - 1]]
            right_card = self.column_frames[remaining[insertion_index]]
            left_edge = left_card.winfo_rootx() + left_card.winfo_width()
            right_edge = right_card.winfo_rootx()
            boundary_root_x = (left_edge + right_edge) // 2

        canvas_left = self.canvas.winfo_rootx() + 2
        canvas_right = canvas_left + max(4, self.canvas.winfo_width() - 4)
        boundary_root_x = max(canvas_left, min(canvas_right, boundary_root_x))
        local_x = boundary_root_x - self.winfo_rootx()
        self.drop_indicator.place(
            x=local_x - self.drop_indicator_width // 2,
            y=self.drop_badge_height - 2,
            width=self.drop_indicator_width,
            height=self.CARD_HEIGHT - self.drop_badge_height + 2,
        )
        badge_x = max(0, min(self.winfo_width() - self.drop_badge_width, local_x - self.drop_badge_width // 2))
        self.drop_badge.place(x=badge_x, y=0)
        self.drop_indicator.lift()
        self.drop_badge.tk.call("raise", self.drop_badge._w)

    def _auto_scroll_during_drag(self, pointer_x: int) -> bool:
        canvas_x = pointer_x - self.canvas.winfo_rootx()
        if canvas_x < 28:
            self.canvas.xview_scroll(-1, "units")
            return True
        elif canvas_x > self.canvas.winfo_width() - 28:
            self.canvas.xview_scroll(1, "units")
            return True
        return False

    def _drag_motion(self, event) -> None:
        if self.drag_key is None:
            return
        # Event coordinates can be several frames old when Windows coalesces a
        # burst of mouse moves, so always query the current system pointer.
        try:
            pointer_x, pointer_y = self.winfo_pointerxy()
        except tk.TclError:
            pointer_x, pointer_y = event.x_root, event.y_root
        if self._auto_scroll_during_drag(pointer_x):
            self.canvas.update_idletasks()
        self.drag_drop_index = self._insertion_index(pointer_x)
        self._show_drop_indicator(self.drag_drop_index)

    def _destroy_drag_visuals(self) -> None:
        root = self.winfo_toplevel()
        if self.drag_motion_bind_id is not None:
            try:
                root.unbind("<B1-Motion>", self.drag_motion_bind_id)
            except tk.TclError:
                pass
            self.drag_motion_bind_id = None
        if self.drag_release_bind_id is not None:
            try:
                root.unbind("<ButtonRelease-1>", self.drag_release_bind_id)
            except tk.TclError:
                pass
            self.drag_release_bind_id = None
        self.drop_indicator.place_forget()
        self.drop_badge.place_forget()

    def _finish_drag(self, event=None) -> None:
        if self.drag_key is None or self.drag_finishing:
            return
        self.drag_finishing = True
        if event is not None:
            try:
                pointer_x, _pointer_y = self.winfo_pointerxy()
            except tk.TclError:
                pointer_x = event.x_root
            self.drag_drop_index = self._insertion_index(pointer_x)
        key = self.drag_key
        insertion_index = self.drag_drop_index if self.drag_drop_index is not None else self.order.index(key)
        remaining = [item for item in self.order if item != key]
        remaining.insert(max(0, min(insertion_index, len(remaining))), key)
        self.order = remaining
        self.drag_key = None
        self.drag_drop_index = None
        # End the grab and drag state before any widget tree is rebuilt.  The
        # old order is then replaced once the ButtonRelease event has unwound,
        # preventing duplicate release callbacks and UI freezes.
        self._destroy_drag_visuals()
        self.after_idle(self._complete_drag_render)

    def _complete_drag_render(self) -> None:
        try:
            self._render_columns()
        finally:
            self.drag_finishing = False

    def set_enabled(self, enabled: bool) -> None:
        self._destroy_drag_visuals()
        self.enabled = bool(enabled)
        self.drag_key = None
        self.drag_drop_index = None
        self._render_columns()

    def reset(self) -> None:
        self._destroy_drag_visuals()
        self.drag_key = None
        self.drag_drop_index = None
        self.order = list(CUSTOMIZABLE_RESULT_COLUMN_KEYS)
        for variable in self.title_vars.values():
            variable.set("")
        self._render_columns()

    def get_configuration(self) -> Tuple[List[str], Dict[str, str]]:
        titles = {
            key: variable.get().strip()
            for key, variable in self.title_vars.items()
            if variable.get().strip()
        }
        return list(self.order), titles

    def destroy(self) -> None:
        self._destroy_drag_visuals()
        super().destroy()


class BomMatcherApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self._install_app_icon()
        self.geometry("1500x820")
        self.minsize(1250, 700)

        self.bom_headers: List[str] = []
        self.library_headers: List[str] = []
        self.bom_items: List[BomItem] = []
        self.materials: List[MaterialItem] = []
        self.material_index: Dict[Tuple[str, str, str], List[int]] = {}
        self.result_rows: List[Dict[str, str]] = []
        self.match_scores: List[int] = []
        self.confirmed_rows = set()
        self.auto_matched_rows = set()
        self.warning_rows = set()
        self.edit_widget: Optional[tk.Entry] = None
        self.app_settings = load_app_settings()
        self.auto_load_library_var = tk.BooleanVar(value=bool(self.app_settings.get("auto_load_last_library", False)))
        self.custom_result_columns_var = tk.BooleanVar(value=bool(self.app_settings.get("custom_result_columns_enabled", False)))
        self.result_column_order = normalize_result_column_order(self.app_settings.get("result_column_order", []))
        self.result_column_titles = normalize_result_column_titles(self.app_settings.get("result_column_titles", {}))
        self.app_settings["result_column_order"] = list(self.result_column_order)
        self.app_settings["result_column_titles"] = dict(self.result_column_titles)
        self.match_summary_var = tk.StringVar(value="匹配成功 0/0，已确认 0/0")
        self.matching_window: Optional[tk.Toplevel] = None
        self.matching_label_var = tk.StringVar(value="匹配中")
        self.loading_base_text = "匹配中"
        self.matching_dot_step = 0
        self.background_executor = ThreadPoolExecutor(max_workers=max(2, min(4, WORKER_COUNT)))
        documents_directory = os.path.join(os.path.expanduser("~"), "Documents")
        self.last_export_directory = documents_directory if os.path.isdir(documents_directory) else os.path.expanduser("~")

        self._setup_style()
        self._build_ui()
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        self.after(500, self.auto_load_last_library)



    def _install_app_icon(self) -> None:
        self._app_icon_photo = tk.PhotoImage(width=32, height=32)
        self._app_icon_photo.put("#1677D2", to=(0, 0, 32, 32))
        for y in (6, 14, 22):
            self._app_icon_photo.put("#FFFFFF", to=(5, y, 10, y + 4))
            self._app_icon_photo.put("#FFFFFF", to=(13, y, 27, y + 4))
        self.iconphoto(True, self._app_icon_photo)

    def _apply_app_icon(self, window: tk.Toplevel) -> None:
        window.iconphoto(False, self._app_icon_photo)

    def _setup_style(self) -> None:
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        default_font = ("Microsoft YaHei UI", 10)
        self.option_add("*Font", default_font)
        style.configure("TFrame", background="#F3F4F6")
        style.configure("TLabel", background="#F3F4F6", foreground="#111827")
        style.configure("TButton", padding=(14, 8), borderwidth=1)
        style.map("TButton", background=[("active", "#E8F1FF")])
        style.configure("Toolbar.TFrame", background="#F3F4F6")
        style.configure("Status.TFrame", background="#F3F4F6")
        style.configure("BlueSummary.TLabel", foreground="#0B63CE")
        style.configure("GreenSummary.TLabel", foreground="#16803A")
        style.configure("SidePage.TFrame", background="#F4F4F4")
        style.configure("SideContent.TFrame", background="#FFFFFF")
        style.configure("SideTitle.TLabel", background="#FFFFFF", font=("Microsoft YaHei UI", 16, "bold"))
        style.configure("SideBody.TLabel", background="#FFFFFF", font=("Microsoft YaHei UI", 11))

    def _build_ui(self) -> None:
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)

        self.content_area = ttk.Frame(self)
        self.content_area.grid(row=0, column=0, sticky="nsew")
        self.content_area.columnconfigure(0, weight=1)
        self.content_area.rowconfigure(0, weight=1)

        self.match_page = ttk.Frame(self.content_area)
        self.preview_page = ttk.Frame(self.content_area, padding=12)
        self.settings_page = ttk.Frame(self.content_area, padding=18)
        for page in [self.match_page, self.preview_page, self.settings_page]:
            page.grid(row=0, column=0, sticky="nsew")
        self.match_page.columnconfigure(0, weight=1)
        self.match_page.rowconfigure(1, weight=1)

        toolbar = ttk.Frame(self.match_page, padding=(12, 10, 12, 8), style="Toolbar.TFrame")
        toolbar.grid(row=0, column=0, sticky="ew")
        toolbar.columnconfigure(6, weight=1, minsize=390)
        toolbar.columnconfigure(7, weight=0)

        ttk.Button(toolbar, text="导入初始 BOM", command=self.load_bom).grid(row=0, column=0, padx=(0, 8))
        ttk.Button(toolbar, text="导入物料库", command=self.load_library).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(toolbar, text="自动匹配", command=self.auto_match).grid(row=0, column=2, padx=(0, 8))
        ttk.Button(toolbar, text="导出 Excel", command=self.export_result).grid(row=0, column=3, padx=(0, 8))
        ttk.Button(toolbar, text="清空", command=self.clear_all).grid(row=0, column=4, padx=(0, 16))

        ttk.Label(toolbar, text="搜索规格").grid(row=0, column=5, sticky="e", padx=(0, 8))
        self.search_var = tk.StringVar()
        self.search_entry = ttk.Entry(toolbar, textvariable=self.search_var, width=48)
        self.search_entry.grid(row=0, column=6, sticky="ew")
        self.search_entry.bind("<KeyRelease>", self.refresh_search_results)
        self.match_summary_box = tk.Frame(toolbar, bg="#EAF3FF", highlightbackground="#4A90E2", highlightthickness=1, padx=10, pady=6)
        self.match_summary_box.grid(row=0, column=7, sticky="e", padx=(12, 0))
        self.match_summary_label = tk.Label(
            self.match_summary_box,
            textvariable=self.match_summary_var,
            bg="#EAF3FF",
            fg="#0B63CE",
            font=("Microsoft YaHei UI", 10, "bold"),
        )
        self.match_summary_label.pack()

        vertical_paned = ttk.PanedWindow(self.match_page, orient=tk.VERTICAL)
        vertical_paned.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 8))

        search_frame = ttk.Frame(vertical_paned, padding=(0, 0, 0, 4))
        search_frame.columnconfigure(0, weight=1)

        self.search_tree = GridTable(search_frame, row_height=24, header_height=22, max_lines=1, visible_rows=3, fill_available_width=True)
        self.search_tree.grid(row=0, column=0, sticky="nsew")
        search_frame.rowconfigure(0, weight=1)
        self.search_tree.on_double_click = lambda _row, _header, _bbox: self.apply_selected_material()
        self.setup_search_tree(["物料编码", "物料名称", "规格", "单位"])

        lower_frame = ttk.Frame(vertical_paned)
        lower_frame.columnconfigure(0, weight=1)
        lower_frame.rowconfigure(0, weight=1)

        paned = ttk.PanedWindow(lower_frame, orient=tk.HORIZONTAL)
        paned.grid(row=0, column=0, sticky="nsew")

        left_frame = ttk.Frame(paned)
        right_frame = ttk.Frame(paned)
        paned.add(left_frame, weight=1)
        paned.add(right_frame, weight=1)

        self._build_tree_panel(left_frame, "初始 BOM", is_result=False)
        self._build_tree_panel(right_frame, "匹配好的规格", is_result=True)

        vertical_paned.add(search_frame, weight=0)
        vertical_paned.add(lower_frame, weight=1)

        self.status_var = tk.StringVar(value="请先导入初始 BOM 和物料库。")
        status_bar = ttk.Frame(self.match_page, padding=(12, 6, 12, 6), style="Status.TFrame")
        status_bar.grid(row=2, column=0, sticky="ew")
        status_bar.columnconfigure(0, weight=1)
        status = ttk.Label(status_bar, textvariable=self.status_var, anchor="w")
        status.grid(row=0, column=0, sticky="ew")
        legend = ttk.Frame(status_bar, padding=(8, 2))
        legend.grid(row=0, column=1, sticky="e", padx=(16, 0))
        self._legend_item(legend, "#44DD8A", "自动匹配").grid(row=0, column=0, padx=(0, 16))
        self._legend_item(legend, "#FF6B6B", "需人工复核").grid(row=0, column=1, padx=(0, 16))
        self._legend_item(legend, "#33AE3B", "已确认").grid(row=0, column=2, padx=(0, 16))
        self._legend_item(legend, "#CCE4FF", "选中").grid(row=0, column=3)
        self._build_preview_page()
        self._build_settings_page()
        self.show_page("匹配")
        self._build_page_tabs(row=1)

    def _legend_item(self, parent, color: str, text: str) -> ttk.Frame:
        frame = ttk.Frame(parent)
        swatch = tk.Frame(frame, width=28, height=17, bg=color, highlightbackground="#6B7280", highlightthickness=1)
        swatch.grid(row=0, column=0, padx=(0, 7))
        swatch.grid_propagate(False)
        ttk.Label(frame, text=text, font=("Microsoft YaHei UI", 11, "bold")).grid(row=0, column=1)
        return frame

    def _build_preview_page(self) -> None:
        for child in self.preview_page.winfo_children():
            child.destroy()
        self.preview_page.columnconfigure(0, weight=1)
        self.preview_page.rowconfigure(2, weight=1)
        ttk.Label(
            self.preview_page, text="BOM 预览", font=("Microsoft YaHei UI", 18, "bold")
        ).grid(row=0, column=0, sticky="w", pady=(4, 8))
        self.preview_status_var = tk.StringVar(value="暂无匹配结果，请先在“匹配”页完成匹配。")
        ttk.Label(self.preview_page, textvariable=self.preview_status_var).grid(
            row=1, column=0, sticky="w", pady=(0, 12)
        )
        # No edit/checkbox callbacks: selection, scrolling and column resizing
        # affect only this view and can never modify the matching results.
        self.preview_tree = GridTable(
            self.preview_page, row_height=32, header_height=30,
            max_lines=2, fill_available_width=True,
        )
        self.preview_tree.grid(row=2, column=0, sticky="nsew")
        self._preview_dirty = True

    def refresh_preview(self) -> None:
        columns, titles = self.active_result_columns()
        visible = [(key, title) for key, title in zip(columns, titles) if key != "确认"]
        self.preview_tree.set_columns(
            [key for key, _ in visible], [title for _, title in visible], redraw=False
        )
        rows = []
        for index, result in enumerate(self.result_rows):
            score = self.match_scores[index] if index < len(self.match_scores) else 0
            rows.append([
                (str(score) if score else "") if key == "备注" else result.get(key, "")
                for key, _ in visible
            ])
        self.preview_tree.set_rows(rows)
        self.preview_status_var.set(
            f"共 {len(rows)} 行 · 只读预览，列名和顺序与匹配结果一致；如需修改，请返回“匹配”页。"
            if rows else "暂无匹配结果，请先在“匹配”页完成匹配。"
        )
        self._preview_dirty = False

    def _build_settings_page(self) -> None:
        for child in self.settings_page.winfo_children():
            child.destroy()
        self.settings_page.columnconfigure(0, weight=1)
        self.settings_page.rowconfigure(0, weight=1)

        self._build_ios_page(
            self.settings_page,
            "设置与帮助",
            [
                ("操作与指导", self._settings_help),
                ("导出规格", self._settings_export_specifications),
                ("启动选项", self._settings_startup_options),
                ("关于", self._settings_about),
            ],
        )

    def _build_ios_page(self, page: ttk.Frame, title: str, items: List[Tuple[str, Callable[[tk.Frame], None]]]) -> None:
        is_settings_page = page is self.settings_page
        is_ipad_page = is_settings_page
        bg = "#F5F5F3" if is_ipad_page else "#F2F2F7"
        page.configure(style="SidePage.TFrame")
        shell = tk.Frame(page, bg=bg)
        shell.grid(row=0, column=0, sticky="nsew")
        shell.columnconfigure(0, weight=0, minsize=315 if is_ipad_page else 270)
        shell.columnconfigure(1, weight=1)
        shell.rowconfigure(0, weight=1)

        sidebar_bg = "#FAFAF8" if is_ipad_page else bg
        sidebar_card = RoundedCard(
            shell,
            bg=sidebar_bg,
            outer_bg=bg,
            radius=30 if is_ipad_page else 24,
            border="#D9D9D4" if is_ipad_page else "#E5E5EA",
            padding=10 if is_ipad_page else 8,
        )
        sidebar_card.grid(
            row=0,
            column=0,
            sticky="nsew",
            padx=(12 if is_ipad_page else 0, 0),
            pady=12 if is_ipad_page else 0,
        )
        sidebar = sidebar_card.body
        sidebar.configure(
            bg=sidebar_bg,
            padx=10 if is_ipad_page else 8,
            pady=2 if is_ipad_page else 8,
        )
        sidebar.columnconfigure(0, weight=1)
        sidebar.rowconfigure(1, weight=1)
        tk.Label(
            sidebar,
            text=title,
            bg=sidebar_bg,
            fg="#1C1C1E",
            font=("Microsoft YaHei UI", 19 if is_ipad_page else 22, "bold"),
            anchor="center" if is_ipad_page else "w",
        ).grid(row=0, column=0, sticky="ew", padx=8, pady=(8 if is_ipad_page else 0, 18 if is_ipad_page else 12))

        menu_card = RoundedCard(
            sidebar,
            bg="#FFFFFF",
            outer_bg=sidebar_bg,
            radius=26 if is_ipad_page else 22,
            border="#E3E3DF",
            padding=8 if is_ipad_page else 6,
        )
        menu_card.grid(row=1, column=0, sticky="new", padx=0)
        menu_card.body.columnconfigure(0, weight=1)
        if is_ipad_page:
            menu_card.canvas.configure(height=max(204, len(items) * 62 + 20))

        content_wrap = tk.Frame(shell, bg=bg, padx=0 if is_ipad_page else 12, pady=12 if is_ipad_page else 8)
        content_wrap.grid(
            row=0,
            column=1,
            sticky="nsew",
            padx=(6, 0) if is_ipad_page else 0,
        )
        content_wrap.columnconfigure(0, weight=1)
        content_wrap.rowconfigure(1, weight=1)
        content_title = tk.Label(
            content_wrap,
            text="",
            bg=bg,
            fg="#1C1C1E",
            font=("Microsoft YaHei UI", 19 if is_ipad_page else 22, "bold"),
            anchor="center" if is_ipad_page else "w",
        )
        content_title.grid(row=0, column=0, sticky="ew", padx=4, pady=(8 if is_ipad_page else 0, 18 if is_ipad_page else 12))
        content_canvas = None
        if is_ipad_page:
            content_canvas = tk.Canvas(content_wrap, bg=bg, highlightthickness=0, bd=0)
            content_canvas.grid(row=1, column=0, sticky="nsew")
            content_scrollbar = MobileVerticalScrollbar(content_canvas, command=content_canvas.yview, bg=bg)
            content_scrollbar.place(relx=1.0, x=-1, rely=0.0, relheight=1.0, anchor="ne")
            content_canvas.configure(yscrollcommand=content_scrollbar.set)
            content_frame = tk.Frame(content_canvas, bg=bg)
            content_window = content_canvas.create_window((0, 0), window=content_frame, anchor="nw")
            content_scroll_enabled = False

            def update_content_scroll_state(_event=None) -> None:
                nonlocal content_scroll_enabled
                viewport_width = max(1, content_canvas.winfo_width())
                viewport_height = max(1, content_canvas.winfo_height())
                content_height = max(1, content_frame.winfo_reqheight())
                content_scroll_enabled = content_height > viewport_height + 3
                region_height = content_height if content_scroll_enabled else viewport_height
                content_canvas.configure(scrollregion=(0, 0, viewport_width, region_height))
                if not content_scroll_enabled:
                    # A Canvas can otherwise move into negative coordinates even when
                    # its content is shorter than the viewport, producing a large blank gap.
                    content_canvas.yview_moveto(0)
                    content_scrollbar.set(0.0, 1.0)

            def resize_content_canvas(event) -> None:
                content_canvas.itemconfigure(content_window, width=event.width)
                update_content_scroll_state()
                content_scrollbar.tk.call("raise", content_scrollbar._w)

            content_canvas.bind("<Configure>", resize_content_canvas)
            content_frame.bind("<Configure>", update_content_scroll_state)
            content_scrollbar.tk.call("raise", content_scrollbar._w)

            def enable_content_wheel(_event=None) -> None:
                def scroll_content(wheel_event) -> str:
                    update_content_scroll_state()
                    if not content_scroll_enabled:
                        content_canvas.yview_moveto(0)
                        return "break"
                    units = -int(wheel_event.delta / 120)
                    if units == 0 and wheel_event.delta:
                        units = -1 if wheel_event.delta > 0 else 1
                    if units:
                        content_canvas.yview_scroll(units, "units")
                    return "break"

                content_canvas.bind_all("<MouseWheel>", scroll_content)

            def disable_content_wheel(_event=None) -> None:
                content_canvas.unbind_all("<MouseWheel>")

            content_canvas.bind("<Enter>", enable_content_wheel)
            content_canvas.bind("<Leave>", disable_content_wheel)
        else:
            content_frame = tk.Frame(content_wrap, bg=bg)
            content_frame.grid(row=1, column=0, sticky="nsew")
        content_frame.columnconfigure(0, weight=1)
        content_frame._ios_settings_style = is_ipad_page

        buttons: List[Any] = []

        def select(index: int) -> None:
            for button_index, button in enumerate(buttons):
                selected = button_index == index
                if is_ipad_page:
                    button.set_selected(selected)
                elif selected:
                    button.configure(bg="#DCEBFF", fg="#0A63CE", font=("Microsoft YaHei UI", 11, "bold"))
                else:
                    button.configure(bg="#FFFFFF", fg="#111111", font=("Microsoft YaHei UI", 11))
            content_title.configure(text=items[index][0])
            for child in content_frame.winfo_children():
                child.destroy()
            items[index][1](content_frame)
            if content_canvas is not None:
                content_frame.update_idletasks()
                content_canvas.yview_moveto(0)
                update_content_scroll_state()
                content_scrollbar.tk.call("raise", content_scrollbar._w)

        navigation_icons = (
            [("?", "#4F86C6"), ("列", "#E69B3A"), ("启", "#6C68A8"), ("钥", "#AF6ACF"), ("i", "#8E8E93")]
            if is_settings_page
            else [("⇩", "#4F86C6"), ("规", "#45A36C")]
        )
        for index, (item_title, _builder) in enumerate(items):
            if is_ipad_page:
                icon_text, icon_color = navigation_icons[index % len(navigation_icons)]
                button = AppleNavItem(
                    menu_card.body,
                    text=item_title,
                    icon_text=icon_text,
                    icon_color=icon_color,
                    command=lambda selected=index: select(selected),
                )
                button.grid(row=index, column=0, sticky="ew", padx=6, pady=(6 if index == 0 else 2, 6 if index == len(items) - 1 else 2))
            else:
                row_frame = tk.Frame(menu_card.body, bg="#FFFFFF")
                row_frame.grid(row=index, column=0, sticky="ew", padx=8, pady=(8 if index == 0 else 0, 8 if index == len(items) - 1 else 0))
                row_frame.columnconfigure(0, weight=1)
                button = tk.Button(
                    row_frame,
                    text=f"{item_title}        >",
                    anchor="w",
                    relief="flat",
                    bg="#FFFFFF",
                    fg="#111111",
                    activebackground="#DCEBFF",
                    activeforeground="#0A63CE",
                    font=("Microsoft YaHei UI", 11),
                    padx=14,
                    pady=10,
                    command=lambda selected=index: select(selected),
                )
                button.grid(row=0, column=0, sticky="ew")
            buttons.append(button)
        if page is self.settings_page:
            self.settings_section_selector = select
        select(0)

    def _content_group(self, parent: tk.Frame, title: str = "") -> tk.Frame:
        ios_style = bool(getattr(parent, "_ios_settings_style", False))
        group = RoundedCard(
            parent,
            bg="#FFFFFF",
            outer_bg="#F5F5F3" if ios_style else "#F2F2F7",
            radius=28 if ios_style else 22,
            border="#E3E3DF" if ios_style else "#E5E5EA",
            padding=18 if ios_style else 14,
        )
        group.grid(sticky="ew", padx=0, pady=(0, 14))
        group.body.columnconfigure(0, weight=1)
        if title:
            tk.Label(group.body, text=title, bg="#FFFFFF", fg="#111111", font=("Microsoft YaHei UI", 12, "bold"), anchor="w").grid(row=0, column=0, sticky="ew", pady=(0, 8))
        return group.body

    def _setting_row(self, parent: tk.Frame, row: int, title: str, value: str = "", control=None) -> None:
        parent.columnconfigure(1, weight=1)
        tk.Label(parent, text=title, bg="#FFFFFF", fg="#111111", font=("Microsoft YaHei UI", 10), anchor="w").grid(row=row, column=0, sticky="w", pady=6)
        if control is not None:
            control.grid(row=row, column=1, sticky="e", padx=(16, 0), pady=4)
        else:
            tk.Label(parent, text=value, bg="#FFFFFF", fg="#6B7280", font=("Microsoft YaHei UI", 10), anchor="e").grid(row=row, column=1, sticky="e", padx=(16, 0), pady=6)

    def _ios_entry(self, parent: tk.Frame, text: str, width: int = 38) -> tk.Entry:
        entry = tk.Entry(
            parent,
            relief="flat",
            bd=0,
            bg="#F2F2F7",
            fg="#1C1C1E",
            insertbackground="#1C1C1E",
            highlightthickness=1,
            highlightbackground="#E1E1E6",
            highlightcolor="#4F86C6",
            font=("Microsoft YaHei UI", 10),
            width=width,
        )
        entry.insert(0, text)
        return entry

    def _fit_content_group(self, group: tk.Frame, minimum: int = 150, extra: int = 40) -> None:
        group.update_idletasks()
        group.master.configure(height=max(minimum, group.winfo_reqheight() + extra))

    def _settings_export_specifications(self, parent: tk.Frame) -> None:
        group = self._content_group(parent, "导出规格")
        enabled = bool(self.custom_result_columns_var.get())
        custom_switch = IosSwitch(
            group,
            variable=self.custom_result_columns_var,
            command=self.on_custom_result_columns_changed,
            width=70,
            height=40,
        )
        self._setting_row(group, 1, "启用自定义规则", control=custom_switch)
        tk.Label(
            group,
            text="第一行是第一页右侧结果表的默认列名；第二行可填写替代名称，留空时继续使用默认名称。按住底部拖动区可调整列顺序。",
            bg="#FFFFFF",
            fg="#6B7280",
            font=("Microsoft YaHei UI", 10),
            anchor="w",
            justify="left",
            wraplength=1200,
        ).grid(row=2, column=0, columnspan=2, sticky="ew", pady=(8, 12))

        self.column_rule_editor = ColumnRuleEditor(
            group,
            self.result_column_order,
            self.result_column_titles,
            enabled=enabled,
        )
        self.column_rule_editor.grid(row=3, column=0, columnspan=2, sticky="ew")

        button_frame = tk.Frame(group, bg="#FFFFFF")
        button_frame.grid(row=4, column=0, columnspan=2, sticky="ew", pady=(14, 0))
        self.column_rule_apply_button = AppleButton(
            button_frame,
            text="完成并应用到第一页",
            command=self.apply_column_rule_editor,
            variant="primary",
            width=190,
            height=42,
            enabled=enabled,
            canvas_bg="#FFFFFF",
        )
        self.column_rule_apply_button.pack(side="left")
        self.column_rule_reset_button = AppleButton(
            button_frame,
            text="恢复默认编辑值",
            command=self.reset_column_rule_editor,
            variant="secondary",
            width=160,
            height=42,
            enabled=enabled,
            canvas_bg="#FFFFFF",
        )
        self.column_rule_reset_button.pack(side="left", padx=(10, 0))
        self.column_rule_status_var = tk.StringVar(
            value="编辑和拖动期间不会刷新第一页；点击“完成并应用到第一页”后只同步一次。"
        )
        tk.Label(
            group,
            textvariable=self.column_rule_status_var,
            bg="#FFFFFF",
            fg="#8E8E93",
            font=("Microsoft YaHei UI", 9),
            anchor="w",
        ).grid(row=5, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        self._fit_content_group(group, minimum=390)

    def on_custom_result_columns_changed(self) -> None:
        enabled = bool(self.custom_result_columns_var.get())
        self.app_settings["custom_result_columns_enabled"] = enabled
        saved = self._save_app_settings_safely()
        if hasattr(self, "column_rule_editor") and self.column_rule_editor.winfo_exists():
            self.column_rule_editor.set_enabled(enabled)
        for button_name in ("column_rule_apply_button", "column_rule_reset_button"):
            button = getattr(self, button_name, None)
            if button is not None and button.winfo_exists():
                button.set_enabled(enabled)
        self.setup_result_tree(redraw=False)
        self.refresh_result_tree()
        message = "自定义列规则已启用，当前已保存规则已同步到第一页。" if enabled else "自定义列规则已关闭，第一页已恢复默认列名和顺序。"
        if not saved:
            message += " 设置未能保存。"
        if hasattr(self, "column_rule_status_var"):
            self.column_rule_status_var.set(message)
        self.status_var.set(message)

    def apply_column_rule_editor(self) -> None:
        if not bool(self.custom_result_columns_var.get()) or not hasattr(self, "column_rule_editor"):
            return
        order, titles = self.column_rule_editor.get_configuration()
        self.result_column_order = normalize_result_column_order(order)
        self.result_column_titles = normalize_result_column_titles(titles)
        self.app_settings["custom_result_columns_enabled"] = True
        self.app_settings["result_column_order"] = list(self.result_column_order)
        self.app_settings["result_column_titles"] = dict(self.result_column_titles)
        saved = self._save_app_settings_safely()
        self.setup_result_tree(redraw=False)
        self.refresh_result_tree()
        message = "列名和列顺序已一次性同步到第一页。" if saved else "列规则已同步到第一页，但未能保存到本机配置。"
        self.column_rule_status_var.set(message)
        self.status_var.set(message)

    def reset_column_rule_editor(self) -> None:
        if not bool(self.custom_result_columns_var.get()) or not hasattr(self, "column_rule_editor"):
            return
        self.column_rule_editor.reset()
        self.column_rule_status_var.set("编辑区已恢复默认值；第一页尚未变化，点击“完成并应用到第一页”后生效。")

    def _settings_startup_options(self, parent: tk.Frame) -> None:
        group = self._content_group(parent, "启动选项")
        auto_load_switch = IosSwitch(group, variable=self.auto_load_library_var, command=self.on_auto_load_library_changed)
        self._setting_row(group, 1, "自动加载上次的物料库", control=auto_load_switch)
        last_path = str(self.app_settings.get("last_library_path", ""))
        self._setting_row(group, 2, "上次使用的物料库", os.path.basename(last_path) if last_path else "尚未记录")
        self._setting_row(group, 3, "启动后停留页面", "匹配")
        self._setting_row(group, 4, "导入完成提示", "开启")
        tk.Label(
            group,
            text="开启后，下次启动会自动加载上次的物料库；文件移动或删除时会静默跳过。",
            bg="#FFFFFF",
            fg="#8E8E93",
            font=("Microsoft YaHei UI", 9),
            anchor="w",
            justify="left",
            wraplength=780,
        ).grid(row=5, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        self._fit_content_group(group, minimum=260)

    def _settings_help(self, parent: tk.Frame) -> None:
        group = self._content_group(parent, "操作与指导")
        group.columnconfigure(0, weight=1, uniform="help")
        group.columnconfigure(1, weight=1, uniform="help")

        flow = "推荐流程：导入初始 BOM  →  导入物料库  →  自动匹配  →  人工核对与确认  →  导出 Excel"
        tk.Label(
            group,
            text=flow,
            bg="#EAF3FF",
            fg="#0A63CE",
            font=("Microsoft YaHei UI", 10, "bold"),
            anchor="w",
            justify="left",
            padx=12,
            pady=10,
        ).grid(row=1, column=0, columnspan=2, sticky="ew", pady=(0, 12))

        left_sections = [
            (
                "一、首次使用",
                "本开源版安装依赖后即可运行。先准备初始 BOM 和物料库文件，导入并检查列名识别结果，再执行自动匹配。所有表格均在本机处理。",
            ),
            (
                "二、导入初始 BOM",
                "返回“匹配”页，点击“导入初始 BOM”。支持 XLSX、XLSM、XLS 和 CSV。表格中建议包含规格/型号（或 Comment、Value）、数量、位号和封装；程序会自动识别常见的中英文列名。导入后请检查左下方行数和内容。",
            ),
            (
                "三、导入物料库",
                "点击“导入物料库”，选择本公司的物料库文件。建议包含物料编码、物料名称、规格型号和单位。导入后程序会逐行提取并缓存器件类别、型号、封装族、引脚数、本体尺寸及电气参数；下次读取同一份未修改的文件时直接使用缓存，文件内容或解析规则变化后会自动重建。若标题识别不正确，请先在原表中整理清晰的列标题。",
            ),
            (
                "四、自动匹配与颜色说明",
                "两份表格导入完成后点击“自动匹配”。程序会同时处理阻容感、保险丝、二极管、LED、晶体管/MOS、IC、晶振、开关及连接器等器件，统一识别数值、型号和 QFN32-4×4/QFN4×4-32、C0603/SMT0603 等封装写法。电容未写耐压时按 50V，电阻未写精度时按 ±1%；类型、型号、引脚数、尺寸或电气参数存在硬冲突时不会自动套用，基础型号对应多个后缀型号时也会转为人工复核。浅绿色表示自动匹配成功；红色表示候选接近或仍需人工复核；深绿色表示已经确认；浅蓝色表示当前选中行。",
            ),
        ]
        right_sections = [
            (
                "五、人工查找、套用与确认",
                "选中左侧 BOM 行，在顶部“搜索规格”输入型号、参数、封装或物料编码；在搜索结果中双击候选物料即可套用。右侧结果表支持双击单元格修正内容。核对无误后勾选确认；“一键勾选自动匹配项”只确认自动匹配成功的项目。",
            ),
            (
                "六、导出、清空与页面切换",
                "确认结果后点击“导出 Excel”，系统会弹出 Windows 原生“另存为”窗口；在窗口中选择磁盘和文件夹、修改文件名，再点击“保存”。若已有同名文件，Windows 会要求确认是否覆盖；点击“取消”不会导出。“清空”只清除本次已导入及匹配的数据。",
            ),
            (
                "七、预览与设置选项",
                "“预览”页显示匹配后的 BOM，表格只读，不能修改内容；如需修改，请返回“匹配”页，完成后重新进入预览即可查看最新结果。“设置与帮助 → 导出规格”可调整结果表列名与顺序，点击“完成并应用到第一页”后生效；“启动选项”可设置是否自动加载上次的物料库。",
            ),
            (
                "八、常见异常",
                "导入失败时请检查文件是否损坏、加密或列名不清晰。导出失败时请关闭占用目标文件的 Excel 窗口，并检查保存目录权限。未找到合适候选时，先检查物料库是否包含该器件，再使用搜索功能人工核对。",
            ),
            (
                "九、使用注意事项",
                "软件为离线工具，表格在本机处理。导入前建议备份重要文件，并确保 Excel 文件未损坏、未设置打开密码。匹配结果受原始数据完整性影响；物料编码、规格、精度、封装和工序等关键字段应由使用者最终确认。",
            ),
        ]

        def add_sections(column: int, sections: List[Tuple[str, str]]) -> tk.Frame:
            section_wrap = tk.Frame(group, bg="#FFFFFF")
            section_wrap.grid(row=2, column=column, sticky="nsew", padx=((0, 12) if column == 0 else (12, 0)))
            section_wrap.columnconfigure(0, weight=1)
            for row, (title, body) in enumerate(sections):
                tk.Label(
                    section_wrap,
                    text=title,
                    bg="#FFFFFF",
                    fg="#111827",
                    font=("Microsoft YaHei UI", 10, "bold"),
                    anchor="w",
                ).grid(row=row * 2, column=0, sticky="ew", pady=(0 if row == 0 else 9, 2))
                tk.Label(
                    section_wrap,
                    text=body,
                    bg="#FFFFFF",
                    fg="#4B5563",
                    font=("Microsoft YaHei UI", 9),
                    anchor="nw",
                    justify="left",
                    wraplength=470,
                ).grid(row=row * 2 + 1, column=0, sticky="ew")
            return section_wrap

        left_wrap = add_sections(0, left_sections)
        right_wrap = add_sections(1, right_sections)
        group.update_idletasks()
        required_height = max(left_wrap.winfo_reqheight(), right_wrap.winfo_reqheight()) + 96
        # Reserve the lower part of the page for the final section instead of
        # clipping it inside the rounded card. The card still scrolls normally
        # on genuinely short windows.
        required_body_height = group.winfo_reqheight() + 36
        group.master.configure(height=max(680, required_height, required_body_height))


    def _settings_about(self, parent: tk.Frame) -> None:
        group = self._content_group(parent, "关于")
        self._setting_row(group, 1, "软件名称", APP_TITLE)
        self._setting_row(group, 2, "当前版本", "v1.0 开源版 · MIT")
        self._setting_row(group, 3, "用途", "默认 BOM 与物料库自动/人工辅助匹配")
        self._setting_row(group, 4, "作者", APP_AUTHOR)
        self._setting_row(group, 5, "联系作者", f"微信号 {AUTHOR_WECHAT}")
        self._fit_content_group(group, minimum=340)

    def _build_side_page(self, page: ttk.Frame, title: str, items: List[Tuple[str, str]]) -> None:
        page.columnconfigure(0, weight=0)
        page.columnconfigure(1, weight=1)
        page.rowconfigure(0, weight=1)

        sidebar = tk.Frame(page, bg="#EEF1F5", width=220)
        sidebar.grid(row=0, column=0, sticky="ns")
        sidebar.grid_propagate(False)
        content = tk.Frame(page, bg="#FFFFFF")
        content.grid(row=0, column=1, sticky="nsew")
        content.columnconfigure(0, weight=1)

        title_label = tk.Label(content, text=title, bg="#FFFFFF", fg="#111111", font=("Microsoft YaHei UI", 18, "bold"), anchor="w")
        title_label.grid(row=0, column=0, sticky="ew", padx=28, pady=(24, 10))
        body_label = tk.Label(content, text="", bg="#FFFFFF", fg="#333333", font=("Microsoft YaHei UI", 11), justify="left", anchor="nw", wraplength=780)
        body_label.grid(row=1, column=0, sticky="nsew", padx=28, pady=(0, 20))

        buttons = []

        def select(index: int) -> None:
            for button_index, button in enumerate(buttons):
                if button_index == index:
                    button.configure(bg="#DDEBFF", fg="#0B63CE")
                else:
                    button.configure(bg="#EEF1F5", fg="#222222")
            title_label.configure(text=items[index][0])
            body_label.configure(text=items[index][1])

        for index, (item_title, _item_text) in enumerate(items):
            button = tk.Button(
                sidebar,
                text=item_title,
                anchor="w",
                relief="flat",
                bg="#EEF1F5",
                fg="#222222",
                activebackground="#DDEBFF",
                activeforeground="#0B63CE",
                font=("Microsoft YaHei UI", 11),
                padx=16,
                pady=9,
                command=lambda selected=index: select(selected),
            )
            button.grid(row=index, column=0, sticky="ew", padx=10, pady=(8 if index == 0 else 0, 2))
            buttons.append(button)
        sidebar.columnconfigure(0, weight=1)
        select(0)

    def show_page(self, name: str) -> None:
        pages = {
            "匹配": self.match_page,
            "预览": self.preview_page,
            "设置与帮助": self.settings_page,
        }
        if name == "预览" and self._preview_dirty:
            self.refresh_preview()
        self.current_page_name = name
        pages[name].tkraise()
        self._update_page_tabs(name)

    def _update_page_tabs(self, active_name: str) -> None:
        if not hasattr(self, "page_tab_buttons"):
            return
        for name, button in self.page_tab_buttons.items():
            button.set_variant("selected" if name == active_name else "neutral")

    def _build_page_tabs(self, row: int) -> None:
        self.current_page_name = "匹配"
        self.bottom_tabs = ttk.Frame(self, padding=(12, 0, 12, 8), style="Status.TFrame")
        self.bottom_tabs.grid(row=row, column=0, sticky="ew")
        self.page_tab_buttons = {}
        for index, name in enumerate(["匹配", "预览", "设置与帮助"]):
            button = AppleButton(
                self.bottom_tabs,
                text=name,
                variant="selected" if name == "匹配" else "neutral",
                width=126 if name == "设置与帮助" else 88,
                height=44,
                canvas_bg="#F3F4F6",
                command=lambda page_name=name: self.show_page(page_name),
            )
            button.grid(row=0, column=index, sticky="w", padx=(0, 8))
            self.page_tab_buttons[name] = button
        self._update_page_tabs("匹配")

    def on_page_tab_changed(self, _event=None) -> None:
        return

    def _build_tree_panel(self, parent: ttk.Frame, title: str, is_result: bool) -> None:
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(1, weight=1)
        title_bar = ttk.Frame(parent)
        title_bar.grid(row=0, column=0, sticky="ew", pady=(0, 6))
        title_bar.columnconfigure(0, weight=1)
        ttk.Label(title_bar, text=title, font=("Microsoft YaHei UI", 11, "bold")).grid(row=0, column=0, sticky="w")
        if is_result:
            self.confirm_auto_var = tk.BooleanVar(value=False)
            ttk.Checkbutton(title_bar, text="一键勾选自动匹配项", variable=self.confirm_auto_var, command=self.toggle_all_auto_matched).grid(row=0, column=1, sticky="e")

        if is_result:
            tree = GridTable(parent, row_height=24, header_height=22, max_lines=2, fill_available_width=True, on_select=self.on_result_select, on_double_click=self.start_cell_edit, on_checkbox=self.toggle_confirmed_row)
            self.result_tree = tree
            self.setup_result_tree()
        else:
            tree = GridTable(parent, row_height=24, header_height=22, max_lines=2, fill_available_width=True, on_select=self.on_bom_select)
            self.bom_tree = tree
        tree.grid(row=1, column=0, sticky="nsew")

    def setup_result_tree(self, redraw: bool = True) -> None:
        columns, display_columns = self.active_result_columns()
        self.result_tree.set_columns(columns, display_columns, redraw=redraw)
        self._preview_dirty = True

    def active_result_columns(self) -> Tuple[List[str], List[str]]:
        if not bool(self.custom_result_columns_var.get()):
            columns = list(RESULT_COLUMN_KEYS)
            return columns, list(columns)
        customizable_columns = normalize_result_column_order(self.result_column_order)
        columns = ["确认"] + customizable_columns
        display_columns = ["确认"] + [self.result_column_titles.get(column, "").strip() or column for column in customizable_columns]
        return columns, display_columns

    def setup_search_tree(self, headers: List[str]) -> None:
        display_headers = headers or ["物料编码", "物料名称", "规格", "单位"]
        self.search_tree.set_columns(display_headers)

    def load_bom(self) -> None:
        path = filedialog.askopenfilename(
            parent=self,
            title="选择导出的 BOM",
            filetypes=[("表格文件", "*.xlsx *.xlsm *.xls *.csv"), ("所有文件", "*.*")],
        )
        if not path:
            return

        def worker() -> Tuple[List[str], List[BomItem]]:
            headers, rows = read_table(path)
            return headers, parse_bom(headers, rows)

        def on_success(result: Tuple[List[str], List[BomItem]]) -> None:
            headers, bom_items = result
            self.bom_headers = headers
            self.bom_items = bom_items
            self.result_rows = []
            self.match_scores = []
            self.confirmed_rows = set()
            self.auto_matched_rows = set()
            self.warning_rows = set()
            self.sync_confirm_auto_checkbox()
            self.refresh_bom_tree()
            self.refresh_result_tree()
            self.update_match_summary()
            self.status_var.set(f"已导入 BOM：{len(self.bom_items)} 行。请继续导入物料库后点击自动匹配。")

        self.status_var.set("正在导入 BOM，请稍候。")
        self.run_background("导入 BOM", "导入中", worker, on_success, "导入 BOM 失败")

    def load_library(self) -> None:
        path = filedialog.askopenfilename(
            parent=self,
            title="选择物料库",
            filetypes=[("表格文件", "*.xlsx *.xlsm *.xls *.csv"), ("所有文件", "*.*")],
        )
        if not path:
            return

        self._load_library_path(path)

    def _load_library_path(self, path: str, automatic: bool = False) -> None:
        normalized_path = os.path.abspath(path)

        def worker() -> Tuple[List[str], List[MaterialItem], Dict[Tuple[str, str, str], List[int]], int, bool]:
            analysis_cache = load_material_analysis_cache(normalized_path)
            headers, rows = read_table(normalized_path)
            materials = parse_materials_fast(headers, rows, analysis_cache)
            material_index = build_material_index(materials)
            cache_hit_count = sum(1 for material in materials if material.analysis_cache_hit)
            cache_saved = True
            try:
                save_material_analysis_cache(normalized_path, materials)
            except OSError:
                cache_saved = False
            return headers, materials, material_index, cache_hit_count, cache_saved

        def on_success(result: Tuple[List[str], List[MaterialItem], Dict[Tuple[str, str, str], List[int]], int, bool]) -> None:
            headers, materials, material_index, cache_hit_count, cache_saved = result
            self.library_headers = headers
            self.materials = materials
            self.material_index = material_index
            self.setup_search_tree(headers)
            self.refresh_search_results()
            self.update_match_summary()
            self.app_settings["last_library_path"] = normalized_path
            settings_saved = self._save_app_settings_safely()
            action = "已自动加载上次的物料库" if automatic else "已导入物料库"
            suffix = "" if settings_saved else "（上次文件路径未能保存）"
            if materials and cache_hit_count == len(materials):
                analysis_status = "已直接读取规格缓存"
            elif cache_saved:
                analysis_status = "已完成逐行规格提取并建立缓存"
            else:
                analysis_status = "已完成逐行规格提取（缓存文件未能保存）"
            self.status_var.set(f"{action}：{len(self.materials)} 条，{analysis_status}。{suffix}")

        if automatic:
            self.status_var.set("正在自动加载上次的物料库，请稍候。")
            self.run_background("自动加载物料库", "自动加载中", worker, on_success, "自动加载物料库失败")
        else:
            self.status_var.set("正在导入物料库，请稍候。")
            self.run_background("导入物料库", "导入中", worker, on_success, "导入物料库失败")

    def _save_app_settings_safely(self) -> bool:
        try:
            save_app_settings(self.app_settings)
        except OSError:
            return False
        return True

    def on_auto_load_library_changed(self) -> None:
        enabled = bool(self.auto_load_library_var.get())
        self.app_settings["auto_load_last_library"] = enabled
        if not self._save_app_settings_safely():
            self.status_var.set("自动加载设置未能保存，请检查当前用户的配置目录权限。")
            return
        if enabled:
            self.status_var.set("已开启自动加载；下次启动会在文件存在时加载上次的物料库。")
        else:
            self.status_var.set("已关闭自动加载上次的物料库。")

    def auto_load_last_library(self) -> None:
        if self.materials or not bool(self.auto_load_library_var.get()):
            return
        path = str(self.app_settings.get("last_library_path", "")).strip()
        if not path or not os.path.isfile(path):
            return
        self._load_library_path(path, automatic=True)

    def refresh_bom_tree(self) -> None:
        columns = self.bom_headers or ["规格", "数量", "位号"]
        self.bom_tree.set_columns(columns)
        table_rows = []
        for index, bom in enumerate(self.bom_items):
            table_rows.append([bom.original.get(column, "") for column in columns])
        self.bom_tree.set_rows(table_rows, list(range(len(table_rows))))
        self.bom_tree.set_auto_matched_ids(self.auto_matched_rows)
        self.bom_tree.set_warning_ids(self.warning_rows - self.confirmed_rows)
        self.bom_tree.set_confirmed_ids(self.confirmed_rows)

    def refresh_result_tree(self) -> None:
        columns, _display_columns = self.active_result_columns()
        table_rows = []
        for index, row in enumerate(self.result_rows):
            score = self.match_scores[index] if index < len(self.match_scores) else 0
            checked = "1" if index in self.confirmed_rows else ""
            values = []
            for column in columns:
                if column == "确认":
                    values.append(checked)
                elif column == "备注":
                    values.append(str(score) if score else "")
                else:
                    values.append(row.get(column, ""))
            table_rows.append(values)
        self.result_tree.set_rows(table_rows, list(range(len(table_rows))), redraw=False)
        self.result_tree.set_row_states(
            confirmed_ids=self.confirmed_rows,
            auto_matched_ids=self.auto_matched_rows,
            warning_ids=self.warning_rows - self.confirmed_rows,
        )
        self._preview_dirty = True
        if getattr(self, "current_page_name", "匹配") == "预览" and hasattr(self, "preview_tree"):
            self.refresh_preview()
        self.update_match_summary()

    def refresh_search_results(self, _event=None) -> None:
        query = self.search_var.get().strip()
        if not self.materials:
            self.search_tree.set_rows([])
            return
        ranked: List[Tuple[int, int, MaterialItem]] = []
        normalized_query = normalize_text(query)
        for index, material in enumerate(self.materials):
            score = match_score(query, material) if query else 1
            material_text = material.normalized_searchable or normalize_text(material.searchable)
            if not query or score >= 25 or normalized_query in material_text:
                ranked.append((score, index, material))
        ranked.sort(key=lambda item: item[0], reverse=True)
        table_rows = []
        row_ids = []
        for _score, index, material in ranked[:80]:
            if self.library_headers:
                values = [material.source.get(header, "") for header in self.library_headers]
            else:
                values = [material.code, material.name, material.spec, material.unit]
            table_rows.append(values)
            row_ids.append(index)
        self.search_tree.set_rows(table_rows, row_ids)
        self.search_tree.canvas.xview_moveto(0)
        self.search_tree.canvas.yview_moveto(0)
        self.search_tree.selected_index = 0 if table_rows else None
        self.search_tree.redraw()

    def on_bom_select(self, row_index: int) -> None:
        row_id = str(row_index)
        self.bom_tree.selection_set(row_id)
        self.result_tree.selection_set(row_id)
        self.result_tree.scroll_row_to_offset_animated(row_id, self.bom_tree.row_view_offset(row_id))

    def on_result_select(self, row_index: int) -> None:
        row_id = str(row_index)
        self.result_tree.selection_set(row_id)
        self.bom_tree.selection_set(row_id)
        self.bom_tree.scroll_row_to_offset_animated(row_id, self.result_tree.row_view_offset(row_id))

    def sync_lower_table_selection(self, row_index: int) -> None:
        row_id = str(row_index)
        self.bom_tree.align_row_to_standard_position(row_id)
        self.result_tree.align_row_to_standard_position(row_id)

    def toggle_confirmed_row(self, row_index: int) -> None:
        if row_index in self.confirmed_rows:
            self.confirmed_rows.remove(row_index)
        else:
            self.confirmed_rows.add(row_index)
        self.sync_confirm_auto_checkbox()
        self.bom_tree.set_confirmed_ids(self.confirmed_rows)
        self.bom_tree.set_warning_ids(self.warning_rows - self.confirmed_rows)
        self.refresh_result_tree()
        self.result_tree.selection_set(str(row_index))
        self.bom_tree.selection_set(str(row_index))
        self.update_match_summary()

    def toggle_all_auto_matched(self) -> None:
        if self.confirm_auto_var.get():
            self.confirmed_rows.update(self.auto_matched_rows)
        else:
            self.confirmed_rows.difference_update(self.auto_matched_rows)
        self.bom_tree.set_confirmed_ids(self.confirmed_rows)
        self.bom_tree.set_warning_ids(self.warning_rows - self.confirmed_rows)
        self.refresh_result_tree()
        self.update_match_summary()

    def sync_confirm_auto_checkbox(self) -> None:
        if not hasattr(self, "confirm_auto_var"):
            return
        if self.auto_matched_rows and self.auto_matched_rows.issubset(self.confirmed_rows):
            self.confirm_auto_var.set(True)
        else:
            self.confirm_auto_var.set(False)

    def update_match_summary(self) -> None:
        total_count = len(self.bom_items)
        matched_count = len(self.auto_matched_rows)
        confirmed_matched_count = len(self.confirmed_rows & self.auto_matched_rows)
        self.match_summary_var.set(f"匹配成功 {matched_count}/{total_count}，已确认 {confirmed_matched_count}/{matched_count}")
        if matched_count > 0 and confirmed_matched_count >= matched_count:
            self.match_summary_label.configure(fg="#16803A")
            self.match_summary_box.configure(highlightbackground="#33AE3B")
        else:
            self.match_summary_label.configure(fg="#0B63CE")
            self.match_summary_box.configure(highlightbackground="#4A90E2")
        self.sync_confirm_auto_checkbox()

    def auto_match(self) -> None:
        if not self.bom_items:
            messagebox.showinfo(APP_TITLE, "请先导入初始 BOM。", parent=self)
            return
        if not self.materials:
            messagebox.showinfo(APP_TITLE, "请先导入物料库。", parent=self)
            return

        self.result_rows = []
        self.match_scores = []
        self.confirmed_rows = set()
        self.auto_matched_rows = set()
        self.warning_rows = set()
        self.sync_confirm_auto_checkbox()
        self.show_matching_window()
        self.after(80, lambda: self._auto_match_batch(0, 0))

    def run_background(
        self,
        title: str,
        message: str,
        worker: Callable[[], Any],
        on_success: Callable[[Any], None],
        error_title: str,
    ) -> None:
        self.show_matching_window(title, message)
        future = self.background_executor.submit(worker)
        self.after(120, lambda: self._poll_background(future, on_success, error_title))

    def _poll_background(self, future: Future, on_success: Callable[[Any], None], error_title: str) -> None:
        if not future.done():
            self.after(120, lambda: self._poll_background(future, on_success, error_title))
            return
        self.close_matching_window()
        try:
            result = future.result()
        except Exception as exc:
            self.show_error(error_title, exc)
            return
        on_success(result)

    def show_matching_window(self, title: str = "自动匹配", base_text: str = "匹配中") -> None:
        if self.matching_window is not None:
            return
        self.loading_base_text = base_text
        self.matching_dot_step = 0
        self.matching_label_var.set(base_text)
        window = tk.Toplevel(self)
        window.title(title)
        self._apply_app_icon(window)
        window.transient(self)
        window.resizable(False, False)
        window.protocol("WM_DELETE_WINDOW", lambda: None)
        ttk.Label(window, textvariable=self.matching_label_var, padding=(28, 18), font=("Microsoft YaHei UI", 11)).pack()
        window.update_idletasks()
        x = self.winfo_rootx() + (self.winfo_width() - window.winfo_width()) // 2
        y = self.winfo_rooty() + (self.winfo_height() - window.winfo_height()) // 2
        window.geometry(f"+{max(0, x)}+{max(0, y)}")
        self.matching_window = window
        self.animate_matching_window()

    def animate_matching_window(self) -> None:
        if self.matching_window is None:
            return
        dots = "." * (self.matching_dot_step % 4)
        self.matching_label_var.set(f"{self.loading_base_text}{dots}")
        self.matching_dot_step += 1
        self.after(350, self.animate_matching_window)

    def close_matching_window(self) -> None:
        if self.matching_window is not None:
            self.matching_window.destroy()
            self.matching_window = None

    def _auto_match_batch(self, start_index: int, matched_count: int) -> None:
        batch_size = 40
        end_index = min(len(self.bom_items), start_index + batch_size)
        for index in range(start_index, end_index):
            bom = self.bom_items[index]
            material, score, warning = best_material_match(bom, self.materials, self.material_index)
            if material and score >= AUTO_MATCH_THRESHOLD:
                matched_count += 1
                self.auto_matched_rows.add(index)
                if warning:
                    self.warning_rows.add(index)
            self.result_rows.append(build_result_row(bom, material, score))
            self.match_scores.append(score)
        self.matching_label_var.set(f"匹配中... {end_index}/{len(self.bom_items)}")
        if end_index < len(self.bom_items):
            self.after(10, lambda: self._auto_match_batch(end_index, matched_count))
            return

        self.bom_tree.set_auto_matched_ids(self.auto_matched_rows)
        self.bom_tree.set_warning_ids(self.warning_rows - self.confirmed_rows)
        self.bom_tree.set_confirmed_ids(self.confirmed_rows)
        self.bom_tree.redraw()
        self.refresh_result_tree()
        total_count = len(self.bom_items)
        warning_count = len(self.warning_rows - self.confirmed_rows)
        warning_text = f"，其中 {warning_count} 行经二次检查后标记为需复核" if warning_count else ""
        self.status_var.set(f"自动匹配完成：{matched_count}/{total_count} 行通过结构化匹配{warning_text}。未通过行可用顶部搜索手动套用。")
        self.close_matching_window()
        messagebox.showinfo(APP_TITLE, f"自动匹配完成\n\n总数：{total_count}\n自动匹配成功：{matched_count}\n二次检查需复核：{warning_count}\n未匹配成功：{total_count - matched_count}", parent=self)

    def selected_row_index(self) -> Optional[int]:
        selection = self.bom_tree.selection() or self.result_tree.selection()
        if not selection:
            return None
        try:
            return int(selection[0])
        except ValueError:
            return None

    def selected_material(self) -> Optional[MaterialItem]:
        selection = self.search_tree.selection()
        if not selection:
            return None
        try:
            return self.materials[int(selection[0])]
        except (ValueError, IndexError):
            return None

    def apply_selected_material(self) -> None:
        row_index = self.selected_row_index()
        material = self.selected_material()
        if row_index is None:
            messagebox.showinfo(APP_TITLE, "请先在左侧 BOM 或右侧结果中选中一行。", parent=self)
            return
        if material is None:
            messagebox.showinfo(APP_TITLE, "请先在顶部搜索结果中选中一个物料。", parent=self)
            return
        if row_index >= len(self.bom_items):
            return
        if not self.result_rows:
            self.result_rows = [build_result_row(item, None, 0) for item in self.bom_items]
            self.match_scores = [0 for _ in self.bom_items]
        while row_index >= len(self.result_rows):
            self.result_rows.append(build_result_row(self.bom_items[len(self.result_rows)], None, 0))
            self.match_scores.append(0)

        self.result_rows[row_index]["物料编码"] = material.code
        self.result_rows[row_index]["物料名称"] = material.name
        self.result_rows[row_index]["规格"] = material.spec
        self.result_rows[row_index]["单位"] = material.unit or "pcs"
        if row_index < len(self.match_scores):
            self.match_scores[row_index] = 100
        self.auto_matched_rows.add(row_index)
        self.warning_rows.discard(row_index)
        self.refresh_result_tree()
        self.bom_tree.set_auto_matched_ids(self.auto_matched_rows)
        self.bom_tree.set_warning_ids(self.warning_rows - self.confirmed_rows)
        self.update_match_summary()
        self.result_tree.selection_set(str(row_index))
        self.bom_tree.selection_set(str(row_index))
        self.sync_lower_table_selection(row_index)
        self.status_var.set(f"已将物料 {material.code or material.name or material.spec} 套用到第 {row_index + 1} 行。")

    def start_cell_edit(self, visible_row_index: int, column: str, _bbox=None) -> None:
        if self.edit_widget is not None:
            self.edit_widget.destroy()
            self.edit_widget = None

        if column in ["备注", "确认"]:
            return
        if visible_row_index < 0 or visible_row_index >= len(self.result_tree.row_ids):
            return
        row_id = self.result_tree.row_ids[visible_row_index]
        current = self.result_rows[row_id].get(column, "") if row_id < len(self.result_rows) else ""
        entry = ttk.Entry(self.result_tree.canvas)
        entry.insert(0, current)
        entry.select_range(0, tk.END)
        self.result_tree.place_editor(visible_row_index, column, entry)
        entry.focus_set()
        self.edit_widget = entry

        def commit(_event=None):
            try:
                if row_id < len(self.result_rows):
                    self.result_rows[row_id][column] = entry.get()
                    self.refresh_result_tree()
                    self.result_tree.selection_set(str(row_id))
                    self.bom_tree.selection_set(str(row_id))
            finally:
                if entry.winfo_exists():
                    entry.destroy()
                self.edit_widget = None

        entry.bind("<Return>", commit)
        entry.bind("<FocusOut>", commit)
        entry.bind("<Escape>", lambda _event: entry.destroy())

    def export_result(self) -> None:
        if not self.result_rows:
            messagebox.showinfo(APP_TITLE, "没有可导出的匹配结果。", parent=self)
            return

        # tk_getSaveFile is the native Windows Save As dialog. Windows 7 uses
        # its Explorer-style dialog; newer Windows versions follow their own theme.
        path = filedialog.asksaveasfilename(
            parent=self,
            title="导出 Excel - 选择保存路径并重命名",
            initialdir=self.last_export_directory,
            defaultextension=".xlsx",
            initialfile="BOM匹配结果.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
            confirmoverwrite=True,
        )
        if not path:
            self.status_var.set("已取消导出。")
            return
        if os.path.splitext(path)[1].lower() != ".xlsx":
            path = os.path.splitext(path)[0] + ".xlsx"
        self.last_export_directory = os.path.dirname(path) or self.last_export_directory
        try:
            export_xlsx(path, self.result_rows)
            self.status_var.set(f"已导出：{path}")
            messagebox.showinfo(APP_TITLE, f"导出完成。\n\n保存位置：{path}", parent=self)
        except Exception as exc:
            self.show_error("导出失败", exc)

    def clear_all(self) -> None:
        self.bom_headers = []
        self.library_headers = []
        self.bom_items = []
        self.materials = []
        self.material_index = {}
        self.result_rows = []
        self.match_scores = []
        self.confirmed_rows = set()
        self.auto_matched_rows = set()
        self.warning_rows = set()
        self.sync_confirm_auto_checkbox()
        self.search_var.set("")
        self.refresh_bom_tree()
        self.setup_search_tree(["物料编码", "物料名称", "规格", "单位"])
        self.refresh_result_tree()
        self.refresh_search_results()
        self.update_match_summary()
        self.status_var.set("已清空。")









    def on_close(self) -> None:
        try:
            self.background_executor.shutdown(wait=False, cancel_futures=True)
        except Exception:
            pass
        self.destroy()

    def show_error(self, title: str, exc: Exception) -> None:
        details = "".join(traceback.format_exception_only(type(exc), exc)).strip()
        messagebox.showerror(title, details, parent=self)
        self.status_var.set(details)


def main() -> None:
    app = BomMatcherApp()
    app.mainloop()


if __name__ == "__main__":
    main()
